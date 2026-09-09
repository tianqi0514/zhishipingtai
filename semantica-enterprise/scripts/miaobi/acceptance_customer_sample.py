from __future__ import annotations

import argparse
import hashlib
import json
import re
import time
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

try:
    from .demo_client import DemoClient
    from .acceptance_upload_to_report import parse_sse
except ImportError:  # Direct script execution.
    from demo_client import DemoClient
    from acceptance_upload_to_report import parse_sse


SPACE_CODE = "miaobi-customer-blind-acceptance"
PRODUCT_CODE = "miaobi-customer-blind-acceptance-product"
PROJECT_CODE = "miaobi-customer-blind-acceptance"
SCENARIO_CODE = "earthquake-response-plan-customer-acceptance"

PLAN_INPUTS = {
    "route_start": "州级应急物资库",
    "route_target": "积石山震中安置点",
    "route_graph": {
        "州级应急物资库": [
            {"to": "快速通道", "minutes": 10, "risk": 7},
            {"to": "安全通道", "minutes": 24, "risk": 1},
            {"to": "综合通道", "minutes": 17.5, "risk": 3},
        ],
        "快速通道": [{"to": "积石山震中安置点", "minutes": 10, "risk": 7}],
        "安全通道": [{"to": "积石山震中安置点", "minutes": 24, "risk": 1}],
        "综合通道": [{"to": "积石山震中安置点", "minutes": 17.5, "risk": 3}],
    },
    "coordinates": {
        "州级应急物资库": [103.21, 35.60],
        "快速通道": [103.31, 35.62],
        "安全通道": [103.22, 35.68],
        "综合通道": [103.27, 35.65],
        "积石山震中安置点": [103.40, 35.66],
    },
}

INPUT_FILENAMES = (
    "临夏州地震应急预案.docx",
    "需求(1).docx",
    "应急救援保障方案生成智能体项目技术要求-0908-1.docx",
    "积石山县6.2级地震应急处置报告推演数据表.xlsx",
    "linxiaearthquakeemergency.ontology.yaml",
    "linxiaearthquakeresponsereport.spec.yaml",
    "linxiaearthquakeemergency.chunks.yaml",
)

HOLDOUT_FILENAMES = {
    "积石山县6.2级地震抗震救灾应急处置方案（下发件）.docx",
    "积石山县6.2级地震应急处置方案推演依据说明.docx",
    "样稿.pdf",
}

REFERENCE = {
    "final_characters": 10849,
    "final_paragraphs": 156,
    "final_pages": 29,
    "evidence_characters": 16173,
    "evidence_paragraphs": 192,
    "evidence_pages": 23,
    "sample_pdf_pages": 19,
}

CHAPTERS = (
    ("general", "一、总则", "说明编制目的、依据、适用范围和处置原则；正文 500—900 字。"),
    ("situation", "二、震情灾情", "写明发生时间、震中、震级、深度、烈度、影响范围和待核事项；正文 700—1100 字。"),
    ("grading", "三、灾害分级与响应启动", "分别写清灾害分级、分级应对、州级响应、启动权限和提级建议；不得把两个维度混为一谈；正文 700—1100 字。"),
    ("command", "四、组织指挥体系", "写清指挥机构、现场指挥部和工作组职责，尽量明确牵头单位；正文 900—1400 字。"),
    ("actions", "五、应急处置任务", "按搜救、医疗、安置、交通、通信、电力、次生灾害等任务写责任单位、时限、输入资源和完成标准；正文 1400—2200 字。"),
    ("safeguards", "六、应急保障", "覆盖队伍、物资、医疗、交通、通信、电力、资金和社会秩序保障，使用确定性缺口数据；正文 1100—1700 字。"),
    ("reporting", "七、信息报送与发布", "说明速报、续报、统一发布、舆情与数据更新机制；正文 500—900 字。"),
    ("discipline", "八、纪律与安全", "说明统一指挥、现场安全、余震边坡监测、信息纪律和责任要求；正文 400—700 字。"),
    ("appendix", "九、附则", "说明方案效力、动态更新、人工确认、演示数据声明和依据追溯；正文 300—600 字。"),
)

REQUIRED_FACTS = {
    "magnitude": 6.2,
    "population_density": 305.6,
    "rescue_required": 500,
    "rescue_available": 320,
    "trauma_beds_required": 330,
    "county_trauma_beds": 110,
    "callable_trauma_beds": 250,
    "tents_required": 7000,
    "tents_available": 5200,
}

EXPECTED_COMPUTATIONS = {
    "rescue_gap": 180,
    "county_bed_gap": 220,
    "all_area_bed_gap": 80,
    "tents_gap": 1800,
}

FACT_LABELS = {
    "event_name": "事件名称",
    "origin_time": "发生时间",
    "epicenter": "震中",
    "magnitude": "震级",
    "depth_km": "震源深度",
    "max_intensity": "最高烈度",
    "population_density": "人口密度",
    "rescue_required": "搜救人员需求",
    "rescue_available": "可用搜救人员",
    "trauma_beds_required": "创伤床位需求",
    "county_trauma_beds": "县域可用床位",
    "callable_trauma_beds": "可调用床位",
    "tents_required": "帐篷需求",
    "tents_available": "可用帐篷",
}

FACT_UNITS = {
    "magnitude": "级",
    "depth_km": "km",
    "population_density": "人/km²",
    "rescue_required": "人",
    "rescue_available": "人",
    "trauma_beds_required": "张",
    "county_trauma_beds": "张",
    "callable_trauma_beds": "张",
    "tents_required": "顶",
    "tents_available": "顶",
}


@dataclass
class Stage:
    name: str
    started_at: float
    elapsed_seconds: float = 0
    status: str = "running"
    detail: dict[str, Any] | None = None


class Recorder:
    def __init__(self) -> None:
        self.stages: list[Stage] = []

    def run(self, name: str, operation: Callable[[], Any]) -> Any:
        stage = Stage(name=name, started_at=time.monotonic())
        self.stages.append(stage)
        print(f"[START] {name}", flush=True)
        try:
            result = operation()
            stage.status = "succeeded"
            if isinstance(result, dict):
                # Acceptance reports must stay reviewable.  A full API response can
                # contain hundreds of Plate nodes or Session Events, which made the
                # first real report grow to several megabytes and hid the actual
                # result.  Keep only short scalar evidence here; detailed business
                # results are recorded in their dedicated report sections.
                stage.detail = {
                    key: value
                    for key, value in list(result.items())[:12]
                    if isinstance(value, (str, int, float, bool, type(None)))
                    and len(str(value)) <= 300
                }
                for key, value in result.items():
                    if isinstance(value, list):
                        stage.detail.setdefault(f"{key}_count", len(value))
            return result
        except Exception as exc:
            stage.status = "failed"
            stage.detail = {"error_type": type(exc).__name__, "message": str(exc)[:500]}
            raise
        finally:
            stage.elapsed_seconds = round(time.monotonic() - stage.started_at, 3)
            print(f"[END] {name}: {stage.status} ({stage.elapsed_seconds}s)", flush=True)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _canonical_hash(value: Any) -> str:
    raw = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _current_version(api: DemoClient, document: dict[str, Any]) -> dict[str, Any] | None:
    detail = api.get(f"/documents/{document['id']}")
    return next(
        (item for item in detail.get("versions", []) if item["id"] == document.get("current_version_id")),
        None,
    )


def ensure_space(api: DemoClient) -> dict[str, Any]:
    row = api.one("/spaces", "code", SPACE_CODE)
    if row:
        return row
    return api.post(
        "/spaces",
        {
            "code": SPACE_CODE,
            "name": "妙笔·客户样本盲测空间",
            "description": "仅使用客户输入材料验证从接入到正式报告导出的完整链路；不包含客户结果样稿。",
            "enabled": True,
        },
    )


def upload_inputs(api: DemoClient, space: dict[str, Any], materials_dir: Path) -> list[dict[str, Any]]:
    unexpected = HOLDOUT_FILENAMES.intersection(path.name for path in materials_dir.iterdir())
    if unexpected:
        raise RuntimeError(f"盲测输入目录包含结果样稿，已拒绝继续：{sorted(unexpected)}")
    paths = [materials_dir / name for name in INPUT_FILENAMES]
    missing = [path.name for path in paths if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"缺少客户输入材料：{missing}")
    existing: dict[str, list[dict[str, Any]]] = {}
    for row in api.get("/documents", space_id=space["id"]):
        existing.setdefault(row["title"], []).append(row)
    result: list[dict[str, Any]] = []
    for path in paths:
        started = time.monotonic()
        expected_sha = _sha256(path)
        document = None
        version = None
        for candidate in existing.get(path.name, []):
            detail = api.get(f"/documents/{candidate['id']}")
            matched = next((item for item in detail.get("versions", []) if item.get("sha256") == expected_sha), None)
            if matched:
                document, version = candidate, matched
                break
        action = "reused"
        if document is None:
            uploaded = api.upload_file(space["id"], path, mode="both")
            document = uploaded["document"]
            version = _current_version(api, document)
            action = "uploaded"
        else:
            summary = (version or {}).get("parse_summary") or {}
            completed = set(summary.get("knowledge_targets_completed") or [])
            if (version or {}).get("status") != "ready" or summary.get("knowledge_status") != "published":
                parse_job = next(
                    (
                        job for job in api.get("/jobs")
                        if job.get("job_type") == "parse_document"
                        and (job.get("input") or {}).get("version_id") == version["id"]
                    ),
                    None,
                )
                if parse_job and parse_job.get("status") == "failed":
                    parse_job = api.post(f"/jobs/{parse_job['id']}/retry")
                if parse_job and parse_job.get("status") != "succeeded":
                    api.wait_job(parse_job["id"])
                api.wait_knowledge_job(version["id"])
                action = "recovered"
            elif not {"vector", "graph"}.issubset(completed):
                job = api.post(f"/documents/{document['id']}/process", {"mode": "both"})
                api.wait_job(job["id"])
                action = "republished"
            version = _current_version(api, document)
        summary = (version or {}).get("parse_summary") or {}
        completed = sorted(set(summary.get("knowledge_targets_completed") or []))
        result.append(
            {
                "filename": path.name,
                "sha256": expected_sha,
                "document_id": document["id"],
                "version_id": (version or {}).get("id"),
                "status": (version or {}).get("status"),
                "parser": summary.get("parser"),
                "knowledge_status": summary.get("knowledge_status"),
                "knowledge_targets_completed": completed,
                "action": action,
                "elapsed_seconds": round(time.monotonic() - started, 3),
            }
        )
        print(f"  {path.name}: {action}, {result[-1]['elapsed_seconds']}s", flush=True)
    return result


def ensure_product_release(api: DemoClient, space: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    product = api.one("/knowledge-products", "code", PRODUCT_CODE)
    if not product:
        product = api.post(
            "/knowledge-products",
            {
                "code": PRODUCT_CODE,
                "name": "妙笔·客户样本盲测知识产品",
                "status": "active",
                "space_ids": [space["id"]],
            },
        )
    else:
        product = api.put(
            f"/knowledge-products/{product['id']}",
            {"status": "active", "enabled": True, "space_ids": [space["id"]]},
        )
    knowledge = api.get("/knowledge/releases", space_id=space["id"])["knowledge"]
    if not knowledge:
        raise RuntimeError("盲测空间没有已发布知识版本")
    knowledge_id = knowledge[0]["id"]
    release = next(
        (
            item for item in api.get(f"/knowledge-products/{product['id']}/releases")
            if item.get("status") == "published"
            and any(link.get("knowledge_release_id") == knowledge_id for link in item.get("items", []))
        ),
        None,
    )
    if not release:
        release = api.post(
            f"/knowledge-products/{product['id']}/releases",
            {"note": "客户样本盲测：仅含输入材料，不含结果样稿"},
        )
    return product, release


def scenario_contract() -> dict[str, Any]:
    return {
        "input_schema": {
            "type": "object",
            "required": list(REQUIRED_FACTS),
            "properties": {key: {"type": "number"} for key in REQUIRED_FACTS},
        },
        "ontology_mapping": {
            "magnitude": "EarthquakeEvent.magnitude",
            "population_density": "AffectedArea.populationDensity",
            "rescue_required": "ResourceDemand.required",
            "rescue_available": "ResourceInventory.available",
        },
        "rule_set_ids": ["earthquake-grade", "earthquake-response-tasks"],
        "formula_ids": ["resource-gap", "medical-pressure", "route-utility"],
        "tool_ids": ["KnowledgeSearch", "SemanticaReasoning", "DeterministicComputation", "MapRoutingMCP"],
        "chapter_template": {
            "chapters": [
                {"key": key, "title": title, "instruction": instruction}
                for key, title, instruction in CHAPTERS
            ]
        },
        "output_schema": {
            "type": "object",
            "required": ["title", "chapters", "evidence_manifest"],
        },
        "review_rules": {
            "require_citations": True,
            "require_action_owner": True,
            "require_action_deadline": True,
            "reject_stale_blocks": True,
            "reject_unconfirmed_gates": True,
        },
        "decision_gates": [
            {"key": "critical_facts", "name": "关键事实确认", "required": True},
            {"key": "disaster_grade", "name": "灾害等级确认", "required": True},
            {"key": "response_level", "name": "响应等级建议确认", "required": True},
            {"key": "resource_gap", "name": "资源缺口确认", "required": True},
            {"key": "alternative_plan", "name": "备选方案选择", "required": True},
            {"key": "publish", "name": "正式发布确认", "required": True},
        ],
        "comparison_dimensions": [
            {"key": "response_time", "name": "响应时间", "direction": "min"},
            {"key": "safety_risk", "name": "安全风险", "direction": "min"},
            {"key": "unmet_demand", "name": "未满足需求", "direction": "min"},
        ],
        "config": {
            "minimum_plan_count": 2,
            "default_plan_count": 3,
            "blind_acceptance": True,
            "reference_final_pages": REFERENCE["final_pages"],
            "reference_final_characters": REFERENCE["final_characters"],
        },
    }


def ensure_scenario(api: DemoClient) -> dict[str, Any]:
    package = api.one("/writing/scenario-packages", "code", SCENARIO_CODE)
    if not package:
        package = api.post(
            "/writing/scenario-packages",
            {
                "code": SCENARIO_CODE,
                "name": "地震应急处置方案（客户样本盲测）",
                "disaster_type": "earthquake",
                "description": "九章正式报告契约；结果样稿不参与生成。",
                "enabled": True,
            },
        )
    detail = api.get(f"/writing/scenario-packages/{package['id']}")
    if detail.get("current_version_id"):
        return next(item for item in detail["versions"] if item["id"] == detail["current_version_id"])
    return api.post(
        f"/writing/scenario-packages/{package['id']}/versions",
        {**scenario_contract(), "activate": True},
    )


def extract_workbook_facts(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    index_rows = list(workbook["00_数据引用索引"].iter_rows(values_only=True))
    index = {str(row[1]): {"value": row[2], "unit": row[3], "row": number} for number, row in enumerate(index_rows[1:], 2) if row[1]}
    resource_rows = list(workbook["R03_资源需求缺口总表"].iter_rows(values_only=True))
    resources = {str(row[0]): {"required": row[4], "county": row[5], "callable": row[6], "gap": row[7], "row": number} for number, row in enumerate(resource_rows[2:], 3) if row[0]}
    grading_rows = list(workbook["R01_灾害分级与响应推演"].iter_rows(values_only=True))
    density_row = next((number, row) for number, row in enumerate(grading_rows, 1) if row[0] == "判据区域成立性")
    density_match = re.search(r"=\s*([0-9.]+)\s*>\s*200", str(density_row[1][2]))
    if not density_match:
        raise RuntimeError("客户数据表中未找到可核验的人口密度计算")
    values: dict[str, tuple[Any, str, int]] = {
        "event_name": ("积石山县6.2级地震", "00_数据引用索引", 2),
        "origin_time": (index["发生时间"]["value"], "00_数据引用索引", index["发生时间"]["row"]),
        "epicenter": (index["震中"]["value"], "00_数据引用索引", index["震中"]["row"]),
        "magnitude": (index["震级"]["value"], "00_数据引用索引", index["震级"]["row"]),
        "depth_km": (index["震源深度"]["value"], "00_数据引用索引", index["震源深度"]["row"]),
        "max_intensity": (index["最高烈度"]["value"], "00_数据引用索引", index["最高烈度"]["row"]),
        "population_density": (float(density_match.group(1)), "R01_灾害分级与响应推演", density_row[0]),
        "rescue_required": (resources["搜救人员"]["required"], "R03_资源需求缺口总表", resources["搜救人员"]["row"]),
        "rescue_available": (resources["搜救人员"]["county"], "R03_资源需求缺口总表", resources["搜救人员"]["row"]),
        "trauma_beds_required": (resources["创伤床位"]["required"], "R03_资源需求缺口总表", resources["创伤床位"]["row"]),
        "county_trauma_beds": (resources["创伤床位"]["county"], "R03_资源需求缺口总表", resources["创伤床位"]["row"]),
        "callable_trauma_beds": (resources["创伤床位"]["callable"], "R03_资源需求缺口总表", resources["创伤床位"]["row"]),
        "tents_required": (resources["帐篷"]["required"], "R03_资源需求缺口总表", resources["帐篷"]["row"]),
        "tents_available": (resources["帐篷"]["county"], "R03_资源需求缺口总表", resources["帐篷"]["row"]),
    }
    result: dict[str, dict[str, Any]] = {}
    for key, (value, sheet, row) in values.items():
        numeric = key in FACT_UNITS and key not in {"max_intensity"}
        result[key] = {
            "value": {"number": value} if numeric else {"text": str(value)},
            "unit": FACT_UNITS.get(key),
            "sheet": sheet,
            "row": row,
        }
    return result


def ensure_project_facts(
    api: DemoClient,
    project: dict[str, Any],
    workbook_document: dict[str, Any],
    workbook_path: Path,
) -> list[dict[str, Any]]:
    extracted = extract_workbook_facts(workbook_path)
    current = {item["fact_key"]: item for item in api.get(f"/writing/projects/{project['id']}/facts")}
    version = _current_version(api, workbook_document)
    if not version:
        raise RuntimeError("客户数据表没有可追溯文档版本")
    rows = []
    for key, item in extracted.items():
        row = current.get(key)
        if row and (row.get("value") != item["value"] or row.get("unit") != item["unit"]):
            row = None
        if not row:
            row = api.post(
                f"/writing/projects/{project['id']}/facts",
                {
                    "fact_key": key,
                    "label": FACT_LABELS[key],
                    "fact_type": "official_brief" if key in {"event_name", "origin_time", "epicenter", "magnitude", "depth_km", "max_intensity"} else "policy_document",
                    "value": item["value"],
                    "unit": item["unit"],
                    "source_type": "policy_document",
                    "source_id": workbook_document["id"],
                    "source_version": version["id"],
                    "source_locator": {
                        "sheet": item["sheet"],
                        "row": item["row"],
                        "sha256": version.get("sha256"),
                        "extraction": "deterministic-xlsx-cell",
                    },
                    "confidence": 1,
                    "verification_status": "unverified",
                },
            )
        if row.get("verification_status") != "verified":
            row = api.post(
                f"/writing/projects/{project['id']}/facts/{row['id']}/confirm",
                {"decision": "confirm", "reason": "客户样本盲测：管理员逐项核对工作表、行号与数值后确认；不代表正式灾情。"},
            )
        rows.append(row)
    return rows


def ensure_project(
    api: DemoClient,
    release: dict[str, Any],
    scenario: dict[str, Any],
) -> dict[str, Any]:
    project = api.one("/writing/projects", "code", PROJECT_CODE)
    if project:
        if project["knowledge_product_release_id"] != release["id"]:
            project = api.post(
                f"/writing/projects/{project['id']}/knowledge-release",
                {"knowledge_product_release_id": release["id"], "reason": "盲测输入知识版本更新"},
            )
        return project
    application = api.one("/applications", "code", "miaobi-emergency")
    return api.post(
        "/writing/projects",
        {
            "code": PROJECT_CODE,
            "name": "积石山县6.2级地震应急处置方案（客户样本盲测）",
            "application_id": application["id"] if application else None,
            "scenario_package_version_id": scenario["id"],
            "knowledge_product_release_id": release["id"],
            "config": {
                "disclaimer": "质量验收稿，不代表实时灾情或正式指挥决定。",
                "plan_inputs": PLAN_INPUTS,
                "route_coordinates": PLAN_INPUTS["coordinates"],
                "holdout_policy": "客户下发件、推演依据和样稿未参与生成",
            },
        },
    )


def run_reasoning_computation_plans(api: DemoClient, project: dict[str, Any]) -> dict[str, Any]:
    facts = {item["fact_key"]: item for item in api.get(f"/writing/projects/{project['id']}/facts")}
    if not {"criterion_major_magnitude", "criterion_high_population_density"}.issubset(facts):
        api.post(f"/writing/projects/{project['id']}/criteria/evaluate")
    reasoning = api.get(f"/writing/projects/{project['id']}/reasoning-runs")
    if not reasoning:
        api.post(f"/writing/projects/{project['id']}/reason", {"mode": "preview"})
    facts = {item["fact_key"]: item for item in api.get(f"/writing/projects/{project['id']}/facts")}
    if not set(EXPECTED_COMPUTATIONS).issubset(facts):
        api.post(f"/writing/projects/{project['id']}/computations/run-baseline")
    plans = api.get(f"/writing/projects/{project['id']}/plans")
    if len(plans) < 3:
        plans = api.post(f"/writing/projects/{project['id']}/plans/generate", {"count": 3})
    if not any(item.get("status") == "selected" for item in plans):
        selected = next((item for item in plans if item.get("plan_key") == "balanced"), plans[0])
        api.post(
            f"/writing/projects/{project['id']}/plans/{selected['id']}/select",
            {"reason": "客户样本盲测默认采用综合平衡方案；不代表正式指挥决定。"},
        )
    return {
        "facts": api.get(f"/writing/projects/{project['id']}/facts"),
        "reasoning": api.get(f"/writing/projects/{project['id']}/reasoning-runs"),
        "computations": api.get(f"/writing/projects/{project['id']}/computations"),
        "plans": api.get(f"/writing/projects/{project['id']}/plans"),
    }


def _run_turn(api: DemoClient, session_id: str, prompt: str) -> list[tuple[str, dict[str, Any]]]:
    with api.client.stream(
        "POST",
        f"/writing/agent-sessions/{session_id}/messages",
        headers={"Accept": "text/event-stream"},
        json={"content": prompt},
    ) as response:
        DemoClient._raise(response)
        return parse_sse(response)


def generate_sections(api: DemoClient, project: dict[str, Any], document: dict[str, Any]) -> dict[str, Any]:
    session = api.post(
        f"/writing/projects/{project['id']}/agent-sessions",
        {"document_id": document["id"], "start_new": True},
    )
    messages: list[dict[str, Any]] = []
    streamed = 0
    for key, title, instruction in CHAPTERS:
        prompt = (
            f"请撰写《积石山县6.2级地震应急处置方案》中的“{title}”。"
            f"必须先调用 writing_generate_section_draft，section_key={key}；随后按需要调用 knowledge_search 或知识图谱工具。"
            f"{instruction}只输出该章节正文，不要解释工具、内部 ID 或写作过程。"
            "事实与数值必须来自当前任务的已核验事实、确定性测算或锁定知识版本；文档性表述必须使用系统真实引用编号。"
            "把待确认内容明确写成‘待确认’，不得冒充正式指挥决定。"
        )
        before = {
            item["id"] for item in (api.get(f"/writing/agent-sessions/{session['id']}").get("conversation") or {}).get("messages", [])
            if item.get("role") == "assistant"
        }
        streamed += len(_run_turn(api, session["id"], prompt))
        restored = api.get(f"/writing/agent-sessions/{session['id']}")
        candidates = [
            item for item in (restored.get("conversation") or {}).get("messages", [])
            if item.get("role") == "assistant" and item.get("id") not in before
        ]
        if not candidates:
            raise RuntimeError(f"章节 {title} 没有生成持久化回答")
        assistant = candidates[-1]
        if assistant.get("status") != "completed" or not str(assistant.get("content") or "").strip():
            raise RuntimeError(f"章节 {title} 生成失败：{assistant.get('error_message') or assistant.get('status')}")
        messages.append({**assistant, "chapter_key": key, "chapter_title": title})
        print(f"  {title}: {len(str(assistant.get('content') or ''))} 字符", flush=True)
    restored = api.get(f"/writing/agent-sessions/{session['id']}")
    events = (restored.get("conversation") or {}).get("events", [])
    return {"session": restored, "messages": messages, "events": events, "streamed_event_count": streamed}


def generate_report_one_click(
    api: DemoClient,
    project: dict[str, Any],
    document: dict[str, Any],
) -> dict[str, Any]:
    """Exercise the production input -> toolbox -> DSH -> Plate endpoint.

    The older per-section loop remains above as a regression helper, but the
    customer acceptance path must prove the same single action exposed by the
    browser.  The server owns trusted calculations, Semantica conclusions,
    section placement and the final quality gate.
    """
    started = api.post(
        f"/writing/projects/{project['id']}/generate-report",
        {"document_id": document["id"], "title": document["title"]},
    )
    session_id = str((started.get("agent_session") or {}).get("id") or "")
    if not session_id:
        raise RuntimeError("一键生成任务没有创建可恢复的 DSH 写作会话")
    with api.client.stream(
        "POST",
        f"/writing/generation-runs/{started['id']}/agent",
        headers={"Accept": "text/event-stream"},
    ) as response:
        DemoClient._raise(response)
        streamed_events = parse_sse(response)
    completed = api.post(f"/writing/generation-runs/{started['id']}/finalize")
    if completed.get("status") != "completed" or not completed.get("quality_report", {}).get("ok"):
        raise RuntimeError(f"一键生成没有通过生产质量门：{completed.get('quality_report')}")
    restored = api.get(f"/writing/agent-sessions/{session_id}")
    conversation = restored.get("conversation") or {}
    return {
        "run": completed,
        "session": restored,
        "messages": [
            item
            for item in conversation.get("messages") or []
            if item.get("role") == "assistant" and item.get("status") == "completed"
        ],
        "events": conversation.get("events") or [],
        "streamed_event_count": len(streamed_events),
    }


def load_generated_document(
    api: DemoClient,
    document_id: str,
) -> tuple[dict[str, Any], list[dict[str, Any]], dict[int, dict[str, Any]]]:
    document = api.get(f"/writing/documents/{document_id}")
    content = (document.get("current_version") or {}).get("content") or []
    rows = api.get(f"/writing/documents/{document_id}/bindings")
    citations = [item for item in rows if item.get("block_type") == "knowledge_citation"]
    evidence: dict[int, dict[str, Any]] = {}
    for index, item in enumerate(citations, 1):
        metadata = item.get("metadata_json") or {}
        number = int(metadata.get("citation_number") or index)
        evidence[number] = {
            "citation_number": number,
            "query_run_id": item.get("retrieval_query_run_id"),
            "chunk_id": item.get("chunk_id"),
            "rank": metadata.get("rank"),
            "snapshot": {
                "title": metadata.get("source_title"),
                "document_id": item.get("source_id"),
                "version_id": item.get("source_version"),
                "page_number": metadata.get("page_number"),
                "structural_path": metadata.get("structural_path"),
            },
        }
    return document, content, evidence


def _strip_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*|__(.*?)__", lambda match: match.group(1) or match.group(2) or "", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    return text.strip()


def renumber_citations(messages: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[int, dict[str, Any]]]:
    evidence_by_key: dict[tuple[str, str], int] = {}
    evidence: dict[int, dict[str, Any]] = {}
    output = []
    for message in messages:
        local: dict[int, int] = {}
        for citation in message.get("citations") or []:
            key = (str(citation.get("query_run_id") or ""), str(citation.get("chunk_id") or ""))
            if not all(key):
                continue
            number = evidence_by_key.get(key)
            if number is None:
                number = len(evidence_by_key) + 1
                evidence_by_key[key] = number
                evidence[number] = citation
            local[int(citation["citation_number"])] = number
        content = re.sub(
            r"\[(\d+)\]",
            lambda match: f"[{local.get(int(match.group(1)), int(match.group(1)))}]",
            str(message.get("content") or ""),
        )
        output.append({**message, "content": content})
    return output, evidence


def markdown_to_plate(
    messages: list[dict[str, Any]],
    evidence: dict[int, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    nodes: list[dict[str, Any]] = [
        {"id": str(uuid.uuid4()), "type": "h1", "children": [{"text": "积石山县6.2级地震应急处置方案（质量验收稿）"}]},
        {"id": str(uuid.uuid4()), "type": "callout", "children": [{"text": "客户样本盲测生成，不代表实时灾情或正式指挥决定。"}]},
    ]
    bindings: list[dict[str, Any]] = []
    occurrence = 0

    def children_for(text: str) -> list[dict[str, Any]]:
        nonlocal occurrence
        children: list[dict[str, Any]] = []
        cursor = 0
        for match in re.finditer(r"\[(\d+)\]", text):
            number = int(match.group(1))
            citation = evidence.get(number)
            if not citation:
                continue
            if match.start() > cursor:
                children.append({"text": text[cursor:match.start()]})
            occurrence += 1
            snapshot = citation.get("snapshot") or {}
            reference = {
                "id": f"acceptance-citation-{number}-{occurrence}",
                "type": "knowledge_citation",
                "citation_label": f"[{number}]",
                "source_title": snapshot.get("title"),
                "chunk_id": citation.get("chunk_id"),
                "query_run_id": citation.get("query_run_id"),
                "source_id": snapshot.get("document_id"),
                "source_version": snapshot.get("version_id"),
                "source_locator": {
                    "page_number": snapshot.get("page_number"),
                    "structural_path": snapshot.get("structural_path"),
                },
                "freshness_status": "current",
                "children": [{"text": ""}],
            }
            children.append(reference)
            bindings.append({"reference": reference, "citation": citation})
            cursor = match.end()
        if cursor < len(text):
            children.append({"text": text[cursor:]})
        return children or [{"text": ""}]

    for message in messages:
        lines = str(message.get("content") or "").splitlines()
        paragraph: list[str] = []

        def flush() -> None:
            if not paragraph:
                return
            text = _strip_markdown(" ".join(paragraph))
            paragraph.clear()
            if text:
                nodes.append({"id": str(uuid.uuid4()), "type": "p", "children": children_for(text)})

        for raw in lines:
            line = raw.strip()
            if not line or line == "---":
                flush()
                continue
            heading = re.match(r"^(#{1,6})\s+(.*)$", line)
            if heading:
                flush()
                level = min(3, len(heading.group(1)))
                nodes.append({"id": str(uuid.uuid4()), "type": f"h{level}", "children": children_for(_strip_markdown(heading.group(2)))})
                continue
            bullet = re.match(r"^[-*]\s+(.*)$", line)
            numbered = re.match(r"^\d+[.、]\s*(.*)$", line)
            if bullet or numbered:
                flush()
                body = _strip_markdown((bullet or numbered).group(1))
                nodes.append(
                    {
                        "id": str(uuid.uuid4()),
                        "type": "p",
                        "indent": 1,
                        "listStyleType": "disc" if bullet else "decimal",
                        "children": children_for(body),
                    }
                )
                continue
            paragraph.append(line)
        flush()
    nodes.append({"id": str(uuid.uuid4()), "type": "h1", "children": [{"text": "生成依据清单"}]})
    for number, citation in evidence.items():
        snapshot = citation.get("snapshot") or {}
        locator = snapshot.get("structural_path") or (f"第 {snapshot.get('page_number')} 页" if snapshot.get("page_number") else "文档片段")
        nodes.append(
            {
                "id": str(uuid.uuid4()),
                "type": "p",
                "children": [{"text": f"[{number}] {snapshot.get('title') or '未命名来源'} · {locator}"}],
            }
        )
    return nodes, bindings


def persist_document(
    api: DemoClient,
    project: dict[str, Any],
    release: dict[str, Any],
    document: dict[str, Any],
    messages: list[dict[str, Any]],
) -> tuple[dict[str, Any], list[dict[str, Any]], dict[int, dict[str, Any]]]:
    normalized, evidence = renumber_citations(messages)
    content, binding_specs = markdown_to_plate(normalized, evidence)
    for item in binding_specs:
        reference = item["reference"]
        citation = item["citation"]
        snapshot = citation.get("snapshot") or {}
        api.post(
            f"/writing/documents/{document['id']}/bindings",
            {
                "block_id": reference["id"],
                "block_type": "knowledge_citation",
                "source_type": "policy_document",
                "source_id": snapshot.get("document_id"),
                "source_version": snapshot.get("version_id"),
                "knowledge_product_release_id": release["id"],
                "chunk_id": citation.get("chunk_id"),
                "query_run_id": citation.get("query_run_id"),
                "content_hash": _canonical_hash(reference),
                "block_content": reference,
                "verification_status": "verified",
                "freshness_status": "current",
                "metadata": {"inserted_from": "customer_sample_blind_acceptance", "citation_number": reference["citation_label"]},
            },
        )
    api.post(
        f"/writing/documents/{document['id']}/versions",
        {"content": content, "change_summary": "九章客户样本盲测报告写回 Plate", "publish": False},
    )
    validation = api.post(f"/writing/documents/{document['id']}/validate", {"for_publish": False})
    if validation.get("issues"):
        raise RuntimeError(f"写回 Plate 后存在可信绑定问题：{validation['issues']}")
    return api.get(f"/writing/documents/{document['id']}"), content, evidence


def confirm_test_gates(api: DemoClient, project: dict[str, Any]) -> list[dict[str, Any]]:
    gates = api.get(f"/writing/projects/{project['id']}/decision-gates")
    for gate in gates:
        if gate.get("status") == "confirmed":
            continue
        api.post(
            f"/writing/projects/{project['id']}/decision-gates/{gate['id']}/records",
            {
                "decision": "confirm",
                "original_value": {"status": gate.get("status")},
                "new_value": {"status": "confirmed_for_quality_acceptance"},
                "reason": "客户样本盲测由管理员确认导出，仅用于质量验收，不代表正式应急决策。",
            },
        )
    return api.get(f"/writing/projects/{project['id']}/decision-gates")


def export_document(api: DemoClient, document: dict[str, Any], output_dir: Path) -> list[dict[str, Any]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    result = []
    for output_format in ("docx", "evidence_docx", "pdf", "json", "xlsx"):
        job = api.post(f"/writing/documents/{document['id']}/exports", {"output_format": output_format})
        response = api.client.get(f"/writing/exports/{job['id']}/download")
        DemoClient._raise(response)
        filename = str((job.get("manifest") or {}).get("filename") or f"report.{output_format}")
        target = output_dir / filename
        target.write_bytes(response.content)
        result.append(
            {
                "format": output_format,
                "job_id": job["id"],
                "filename": filename,
                "path": str(target),
                "size": target.stat().st_size,
                "checksum": _sha256(target),
            }
        )
    return result


def plain_text(nodes: list[dict[str, Any]]) -> str:
    def node_text(node: dict[str, Any]) -> str:
        if "text" in node:
            return str(node.get("text") or "")
        return "".join(node_text(item) for item in node.get("children") or [] if isinstance(item, dict))

    return "\n".join(node_text(node) for node in nodes)


def _walk_nodes(nodes: list[dict[str, Any]]):
    for node in nodes:
        yield node
        children = node.get("children") or []
        if children:
            yield from _walk_nodes([item for item in children if isinstance(item, dict)])


def content_quality_metrics(nodes: list[dict[str, Any]], exports: list[dict[str, Any]]) -> dict[str, Any]:
    text = plain_text(nodes)

    def node_plain_text(node: dict[str, Any]) -> str:
        if "text" in node:
            return str(node.get("text") or "")
        return "".join(
            node_plain_text(item)
            for item in node.get("children") or []
            if isinstance(item, dict)
        )

    sentences = [
        re.sub(r"\s+", "", item)
        for item in re.split(r"[。！？\n]+", text)
        if len(re.sub(r"\s+", "", item)) >= 18
    ]
    duplicate_sentence_count = sum(count - 1 for count in __import__("collections").Counter(sentences).values() if count > 1)
    private_work_patterns = (
        r"\bI have\b",
        r"\bLet me\b",
        r"\bThe (?:section|fragment|project context)\b",
        r"\bVerified facts\b",
        r"\bgraph query returned\b",
        r"\bsearch results\b",
        r"Let me present this as the final answer",
    )
    private_work_matches = sum(len(re.findall(pattern, text, re.IGNORECASE)) for pattern in private_work_patterns)
    platform_process_matches = sum(
        text.count(value)
        for value in (
            "图谱查询返回为空",
            "本轮已调用知识图谱工具",
            "尚未写入文稿",
            "Let me count characters",
            "Citations used",
        )
    )
    walked_nodes = list(_walk_nodes(nodes))
    heading_nodes = [
        node
        for node in walked_nodes
        if str(node.get("type") or "") in {"h1", "h2", "h3"}
    ]
    chapter_occurrences = {
        title: sum(node_plain_text(node).strip() == title for node in heading_nodes)
        for _, title, _ in CHAPTERS
    }
    node_types = [str(node.get("type") or "") for node in walked_nodes]
    docx_row = next((item for item in exports if item.get("format") == "docx"), None)
    docx_metrics: dict[str, Any] = {
        "page_size": None,
        "is_a4": False,
        "visible_inline_citations": 0,
    }
    if docx_row and Path(str(docx_row.get("path") or "")).is_file():
        from docx import Document

        exported = Document(str(docx_row["path"]))
        sizes = [
            (round(float(section.page_width.inches), 2), round(float(section.page_height.inches), 2))
            for section in exported.sections
        ]
        body_paragraphs: list[str] = []
        for paragraph in exported.paragraphs:
            if paragraph.text.strip() == "生成依据清单":
                break
            body_paragraphs.append(paragraph.text)
        body_text = "\n".join(body_paragraphs)
        docx_metrics = {
            "page_size": sizes,
            "is_a4": bool(sizes) and all(abs(width - 8.27) <= 0.05 and abs(height - 11.69) <= 0.05 for width, height in sizes),
            "visible_inline_citations": len(re.findall(r"\[\d+\]", body_text)),
        }
    return {
        "private_work_matches": private_work_matches,
        "platform_process_matches": platform_process_matches,
        "duplicate_sentence_count": duplicate_sentence_count,
        "duplicate_sentence_ratio": round(duplicate_sentence_count / max(1, len(sentences)), 3),
        "chapter_occurrences": chapter_occurrences,
        "computed_metric_nodes": node_types.count("computed_metric"),
        "inference_conclusion_nodes": node_types.count("inference_conclusion"),
        "knowledge_citation_nodes": node_types.count("knowledge_citation"),
        "docx": docx_metrics,
    }


def assess(
    *,
    uploads: list[dict[str, Any]],
    extracted_facts: list[dict[str, Any]],
    execution: dict[str, Any],
    generated: dict[str, Any],
    content: list[dict[str, Any]],
    evidence: dict[int, dict[str, Any]],
    exports: list[dict[str, Any]],
    stages: list[Stage],
) -> dict[str, Any]:
    text = plain_text(content)
    content_metrics = content_quality_metrics(content, exports)
    facts = {item["fact_key"]: item for item in extracted_facts}
    current_facts = {item["fact_key"]: item for item in execution["facts"]}
    computation_values = {
        key: (current_facts.get(key, {}).get("value") or {}).get("number")
        for key in EXPECTED_COMPUTATIONS
    }
    owner_count = len(re.findall(r"(?:牵头|负责|责任单位)", text))
    deadline_count = len(re.findall(r"\d+\s*(?:分钟|小时|日)", text))
    tools = {
        str((item.get("payload") or {}).get("name") or "")
        for item in generated["events"]
        if item.get("event_type") == "tool_finished" and (item.get("payload") or {}).get("success") is not False
    }
    tools.discard("")
    checks = [
        {
            "name": "结果样稿未进入生成输入",
            "passed": True,
            "evidence": f"输入严格限定为 {len(INPUT_FILENAMES)} 份源材料",
        },
        {
            "name": "源材料全部解析并发布向量与图谱",
            "passed": len(uploads) == len(INPUT_FILENAMES) and all(
                row.get("status") == "ready"
                and row.get("knowledge_status") == "published"
                and {"vector", "graph"}.issubset(row.get("knowledge_targets_completed") or [])
                for row in uploads
            ),
            "evidence": f"{sum(1 for row in uploads if row.get('status') == 'ready')}/{len(uploads)} 份 ready",
        },
        {
            "name": "客户数据表事实按工作表和行号进入任务",
            "passed": all(
                (facts.get(key, {}).get("value") or {}).get("number") == value
                for key, value in REQUIRED_FACTS.items()
            ),
            "evidence": f"{len(facts)} 项事实带 spreadsheet locator",
        },
        {
            "name": "Semantica 规则推演成功",
            "passed": any(item.get("status") == "succeeded" and item.get("engine") == "semantica-datalog" for item in execution["reasoning"]),
            "evidence": f"{len(execution['reasoning'])} 次推演运行",
        },
        {
            "name": "确定性资源缺口与 Ground Truth 一致",
            "passed": computation_values == EXPECTED_COMPUTATIONS,
            "evidence": json.dumps(computation_values, ensure_ascii=False),
        },
        {
            "name": "三套算法方案真实不同",
            "passed": len(execution["plans"]) >= 3 and len({json.dumps(item.get("result"), sort_keys=True) for item in execution["plans"]}) >= 3,
            "evidence": f"{len(execution['plans'])} 套方案，{len({json.dumps(item.get('result'), sort_keys=True) for item in execution['plans']})} 种结果",
        },
        {
            "name": "DSH 一次完成九章写作并调用知识工具",
            "passed": len(generated["messages"]) == 1 and {"writing_get_project_context", "knowledge_search"}.issubset(tools),
            "evidence": "、".join(sorted(tools)),
        },
        {
            "name": "九章正式结构完整",
            "passed": all(title in text for _, title, _ in CHAPTERS),
            "evidence": f"命中 {sum(title in text for _, title, _ in CHAPTERS)}/{len(CHAPTERS)} 章",
        },
        {
            "name": "正文体量处于客户样本的 75%—140%",
            "passed": int(REFERENCE["final_characters"] * 0.75) <= len(text) <= int(REFERENCE["final_characters"] * 1.4),
            "evidence": f"生成 {len(text)} 字符 / 客户下发件 {REFERENCE['final_characters']} 字符",
        },
        {
            "name": "责任单位与时限具有可执行性",
            "passed": owner_count >= 12 and deadline_count >= 6,
            "evidence": f"责任表达 {owner_count} 处，时限 {deadline_count} 处",
        },
        {
            "name": "正文引用达到多来源可核验门槛",
            "passed": len(evidence) >= 10 and len({(item.get("snapshot") or {}).get("document_id") for item in evidence.values()}) >= 3,
            "evidence": f"{len(evidence)} 条引用，{len({(item.get('snapshot') or {}).get('document_id') for item in evidence.values()})} 份来源",
        },
        {
            "name": "Plate 原生节点与引用绑定已保存",
            "passed": bool(content) and not any("<" in str(node.get("text") or "") for node in content),
            "evidence": f"{len(content)} 个顶层 Plate 节点",
        },
        {
            "name": "正式稿与独立依据报告均真实导出",
            "passed": {item["format"] for item in exports} == {"docx", "evidence_docx", "pdf", "json", "xlsx"} and all(item["size"] > 0 for item in exports),
            "evidence": "、".join(f"{item['format']}={item['size']}B" for item in exports),
        },
        {
            "name": "正式正文没有模型工作草稿或内部执行说明",
            "passed": content_metrics["private_work_matches"] == 0 and content_metrics["platform_process_matches"] == 0,
            "evidence": f"英文工作草稿 {content_metrics['private_work_matches']} 处，平台过程说明 {content_metrics['platform_process_matches']} 处",
        },
        {
            "name": "每个一级章节只出现一次",
            "passed": all(count == 1 for count in content_metrics["chapter_occurrences"].values()),
            "evidence": json.dumps(content_metrics["chapter_occurrences"], ensure_ascii=False),
        },
        {
            "name": "正文不存在明显重复生成",
            "passed": content_metrics["duplicate_sentence_ratio"] <= 0.05,
            "evidence": f"重复长句 {content_metrics['duplicate_sentence_count']} 条，占比 {content_metrics['duplicate_sentence_ratio']:.1%}",
        },
        {
            "name": "正式 DOCX 使用 A4 页面",
            "passed": content_metrics["docx"]["is_a4"],
            "evidence": f"页面尺寸 {content_metrics['docx']['page_size']}",
        },
        {
            "name": "DOCX 正文保留可见引用编号",
            "passed": content_metrics["docx"]["visible_inline_citations"] >= 10,
            "evidence": f"依据清单前可见引用 {content_metrics['docx']['visible_inline_citations']} 个",
        },
        {
            "name": "确定性测算和推演结论使用可信业务块",
            "passed": content_metrics["computed_metric_nodes"] >= len(EXPECTED_COMPUTATIONS)
            and content_metrics["inference_conclusion_nodes"] >= 1,
            "evidence": f"测算块 {content_metrics['computed_metric_nodes']}，推演块 {content_metrics['inference_conclusion_nodes']}",
        },
    ]
    passed = sum(1 for item in checks if item["passed"])
    score = round(100 * passed / len(checks))
    generation_seconds = round(sum(stage.elapsed_seconds for stage in stages if stage.name == "一键生成九章报告"), 3)
    gaps = []
    if len(text) < int(REFERENCE["final_characters"] * 0.75):
        gaps.append("生成正文低于客户下发件的 75%，需加强章节最小信息契约和自动续写。")
    if len(evidence) < 10:
        gaps.append("引用密度不足，写作编排应按章节强制检索并校验引用覆盖。")
    if not all(title in text for _, title, _ in CHAPTERS):
        gaps.append("场景包没有稳定约束九章结构，需在服务端做章节覆盖校验。")
    if content_metrics["private_work_matches"] or content_metrics["platform_process_matches"]:
        gaps.append("模型工作草稿和工具执行说明混入正式正文；服务端必须在写回前执行结构化产物校验，命中即阻断而不是靠提示词约束。")
    if any(count != 1 for count in content_metrics["chapter_occurrences"].values()):
        gaps.append("多轮回答中的初稿与最终稿被一起写回，造成一级章节重复；写作协议需要单独的 structured_final_document 产物，不能直接拼接 Assistant 全量文本。")
    if content_metrics["duplicate_sentence_ratio"] > 0.05:
        gaps.append("正文存在明显重复生成，需增加章节级去重、字数上限和跨章节一致性审校。")
    if not content_metrics["docx"]["is_a4"]:
        gaps.append("正式导出使用 Letter 页面且未真正应用客户模板，需改为 A4 公文模板并验证分页、文号、版记和字体。")
    if content_metrics["docx"]["visible_inline_citations"] < 10:
        gaps.append("Plate 中引用可点击，但 DOCX 正文未渲染可见引用编号；导出器必须把 knowledge_citation 节点渲染为可核验脚注或尾注。")
    if content_metrics["computed_metric_nodes"] < len(EXPECTED_COMPUTATIONS) or content_metrics["inference_conclusion_nodes"] < 1:
        gaps.append("数值和推演结论仍以普通文字写入正文，没有绑定 ComputationRun/InferenceRun 可信业务块，无法进行局部失效和重算。")
    if not generated.get("run", {}).get("quality_report", {}).get("ok"):
        gaps.append("一键生成结果没有通过服务端生产质量门。")
    if "evidence_docx" not in {item["format"] for item in exports}:
        gaps.append("正式方案缺少独立的生成依据报告。")
    return {
        "score": score,
        "verdict": "达到客户样本基线" if score >= 92 and not gaps else "未达到客户样本基线",
        "checks": checks,
        "failed_checks": [item["name"] for item in checks if not item["passed"]],
        "generated_characters": len(text),
        "reference_characters": REFERENCE["final_characters"],
        "length_ratio": round(len(text) / REFERENCE["final_characters"], 3),
        "citation_count": len(evidence),
        "source_document_count": len({(item.get("snapshot") or {}).get("document_id") for item in evidence.values()}),
        "successful_tools": sorted(tools),
        "generation_seconds": generation_seconds,
        "content_quality_metrics": content_metrics,
        "gaps": gaps,
    }


def write_markdown_report(path: Path, payload: dict[str, Any]) -> None:
    quality = payload["quality"]
    lines = [
        "# 妙笔客户样本盲测报告",
        "",
        f"- 测试时间：{payload['tested_at']}",
        f"- 结论：{quality['verdict']}",
        f"- 严格评分：{quality['score']}/100",
        f"- 正文体量：{quality['generated_characters']} / {quality['reference_characters']} 字符（{quality['length_ratio']:.1%}）",
        f"- 引用：{quality['citation_count']} 条，覆盖 {quality['source_document_count']} 份材料",
        f"- 九章生成耗时：{quality['generation_seconds']} 秒",
        "",
        "## 验收项",
        "",
        "| 验收项 | 结果 | 证据 |",
        "|---|---|---|",
    ]
    for item in quality["checks"]:
        lines.append(f"| {item['name']} | {'通过' if item['passed'] else '失败'} | {str(item['evidence']).replace('|', '／')} |")
    lines.extend(["", "## 与客户样本的关键差距", ""])
    for index, gap in enumerate(quality["gaps"], 1):
        lines.append(f"{index}. {gap}")
    lines.extend(["", "## 阶段耗时", "", "| 阶段 | 状态 | 秒 |", "|---|---|---:|"])
    for stage in payload["stages"]:
        lines.append(f"| {stage['name']} | {stage['status']} | {stage['elapsed_seconds']} |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="使用客户真实输入执行妙笔盲测，不把结果样稿提供给模型")
    parser.add_argument("--materials-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    materials_dir = args.materials_dir.resolve()
    output_dir = args.output_dir.resolve()
    recorder = Recorder()
    api = DemoClient()

    space = recorder.run("创建隔离知识空间", lambda: ensure_space(api))
    uploads = recorder.run("上传、解析并发布七份真实输入材料", lambda: upload_inputs(api, space, materials_dir))
    product, release = recorder.run("锁定知识产品版本", lambda: ensure_product_release(api, space))
    scenario = recorder.run("激活九章客户验收场景包", lambda: ensure_scenario(api))
    project = recorder.run("创建客户样本盲测任务", lambda: ensure_project(api, release, scenario))
    workbook_document = next(item for item in api.get("/documents", space_id=space["id"]) if item["title"] == "积石山县6.2级地震应急处置报告推演数据表.xlsx")
    facts = recorder.run(
        "从已上传工作簿确定性提取并确认事实",
        lambda: ensure_project_facts(
            api,
            project,
            workbook_document,
            materials_dir / "积石山县6.2级地震应急处置报告推演数据表.xlsx",
        ),
    )
    document = api.post(
        "/writing/documents",
        {
            "project_id": project["id"],
            "title": f"积石山县6.2级地震应急处置方案（盲测-{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')}）",
            "content": [],
        },
    )
    generated = recorder.run("一键生成九章报告", lambda: generate_report_one_click(api, project, document))
    execution = recorder.run("核验推演工具箱结果", lambda: {
        "facts": api.get(f"/writing/projects/{project['id']}/facts"),
        "reasoning": api.get(f"/writing/projects/{project['id']}/reasoning-runs"),
        "computations": api.get(f"/writing/projects/{project['id']}/computations"),
        "plans": api.get(f"/writing/projects/{project['id']}/plans"),
    })
    document, content, evidence = recorder.run("核验 Plate 原生节点和依据绑定", lambda: load_generated_document(api, document["id"]))
    recorder.run("确认验收导出闸门", lambda: confirm_test_gates(api, project))
    exports = recorder.run("真实导出 DOCX/PDF/JSON/XLSX", lambda: export_document(api, document, output_dir))
    quality = assess(
        uploads=uploads,
        extracted_facts=facts,
        execution=execution,
        generated=generated,
        content=content,
        evidence=evidence,
        exports=exports,
        stages=recorder.stages,
    )
    payload = {
        "tested_at": datetime.now(timezone.utc).isoformat(),
        "blind_test": True,
        "holdouts_excluded": sorted(HOLDOUT_FILENAMES),
        "reference_baseline": REFERENCE,
        "space": {"id": space["id"], "code": space["code"], "name": space["name"]},
        "knowledge_product": {"id": product["id"], "code": product["code"], "release": release["version"]},
        "project": {"id": project["id"], "code": project["code"], "name": project["name"]},
        "document": {"id": document["id"], "title": document["title"], "version": (document.get("current_version") or {}).get("version")},
        "uploads": uploads,
        "quality": quality,
        "exports": exports,
        "stages": [stage.__dict__ for stage in recorder.stages],
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "customer-sample-writing-acceptance.json"
    md_path = output_dir / "customer-sample-writing-acceptance.md"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    write_markdown_report(md_path, payload)
    print(
        json.dumps(
            {
                "score": quality["score"],
                "verdict": quality["verdict"],
                "failed_checks": quality["failed_checks"],
                "report": str(md_path),
                "exports": [item["path"] for item in exports],
            },
            ensure_ascii=False,
            indent=2,
        ),
        flush=True,
    )
    if quality["verdict"] != "达到客户样本基线":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
