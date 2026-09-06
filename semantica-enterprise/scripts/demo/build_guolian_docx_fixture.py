#!/usr/bin/env python3
"""Build the deterministic DOCX used by the Guolian customer demo."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "demo" / "guolian" / "智慧流程中枢项目周报（演示版）.docx"
NOTICE = "演示数据，不代表国联集团真实经营数据。"


def _font(run, name: str = "Arial Unicode MS", size: float = 10.5, *, bold: bool = False, color: str = "1F2937"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def _set_cell(cell, text: str, *, header: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    _font(run, size=9.5, bold=header, color="FFFFFF" if header else "1F2937")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        _shade(cell, "2456A6")


def _heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = f"Heading {level}"
    run = paragraph.add_run(text)
    _font(run, size=16 if level == 1 else 13, bold=True, color="173E75")


def _body(document: Document, text: str, *, bullet: bool = False) -> None:
    paragraph = document.add_paragraph(style="List Bullet" if bullet else None)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    _font(run)


def build() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(header.add_run("传神智库 · 客户演示材料"), size=9, color="64748B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(footer.add_run(NOTICE), size=8.5, color="9A3412")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(title.add_run("智慧流程中枢项目周报（演示版）"), size=24, bold=True, color="173E75")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(subtitle.add_run("报告周期：2026 年 4 月 27 日—5 月 3 日"), size=11, color="475569")
    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(notice.add_run(NOTICE), size=10, bold=True, color="B45309")

    _heading(document, "一、本周总体情况")
    _body(document, "项目整体进度 68%，需求梳理、流程建模和知识服务接口已完成阶段验收；NexusOne 部署处于集成联调阶段。")
    _body(document, "东方智造负责供应 NexusOne 部署所需的演示组件。供应商于 2026 年 4 月 28 日提示关键组件预计延期十个工作日。")
    _body(document, "项目管理办公室将该事项标记为高风险，要求数字化管理部牵头形成替代方案，风险管理部持续跟踪。")

    _heading(document, "二、里程碑与责任分工")
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["里程碑", "计划日期", "当前状态", "责任部门", "本周结论"]):
        _set_cell(cell, value, header=True)
    rows = [
        ("制度与流程梳理", "2026-03-31", "已完成", "集团采购管理部", "形成 18 条关键控制规则"),
        ("核心流程配置", "2026-04-30", "已完成", "数字化管理部", "采购、合同和用印流程已联调"),
        ("NexusOne 集成", "2026-05-20", "有风险", "项目管理办公室", "受供应商交付延期影响"),
        ("试运行", "2026-06-15", "待开始", "数字科技公司", "需要在 5 月 15 日前确认替代方案"),
    ]
    for values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value)

    document.add_section(WD_SECTION.NEW_PAGE)
    _heading(document, "三、风险与行动项")
    risk = document.add_table(rows=1, cols=6)
    risk.style = "Table Grid"
    risk.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(risk.rows[0].cells, ["风险", "等级", "影响对象", "责任部门", "完成时限", "处置要求"]):
        _set_cell(cell, value, header=True)
    risk_rows = [
        ("东方智造交付延期", "高", "NexusOne、智慧流程中枢项目", "数字化管理部", "2026-05-15", "确认替代供应与分批交付方案"),
        ("统一身份组件维护窗口", "中", "集团数据交换平台、智慧流程中枢", "数字化管理部", "2026-09-10", "完成维护影响评估与回退预案"),
        ("采购审批记录待补", "中", "DEMO-PO-2026-002", "集团采购管理部", "2026-05-08", "补齐审批依据并形成审计记录"),
    ]
    for values in risk_rows:
        row = risk.add_row()
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value)

    _heading(document, "四、跨系统依赖")
    for item in (
        "智慧流程中枢依赖集团数据交换平台完成采购、合同和审批数据交换。",
        "集团数据交换平台使用统一身份组件进行用户身份校验。",
        "NexusOne 为智慧流程中枢提供语义建模、知识检索和智能问答能力。",
        "上述关系应进入知识图谱，并保留本周报作为来源证据。",
    ):
        _body(document, item, bullet=True)

    _heading(document, "五、下周计划")
    for item in (
        "数字化管理部在 5 月 8 日前取得东方智造分批交付承诺。",
        "项目管理办公室在 5 月 12 日前完成 NexusOne 替代方案评估。",
        "风险管理部每周更新供应商风险状态，并检查其影响路径。",
        "采购管理部补齐已执行订单的审批记录，确保采购流程可追溯。",
    ):
        _body(document, item, bullet=True)

    _heading(document, "附录：可核验事实")
    _body(document, "事实 A：东方智造（演示）供应 NexusOne。")
    _body(document, "事实 B：NexusOne 用于智慧流程中枢项目（演示）。")
    _body(document, "事实 C：智慧流程中枢项目由数字科技公司（演示）的项目管理办公室负责。")
    _body(document, "事实 D：东方智造（演示）存在交付延期风险。")
    _body(document, "这些事实用于演示全文、向量、图谱检索和可追溯规则推演，不能视为真实集团经营信息。")

    for style_name in ("Normal", "Body Text"):
        style = document.styles[style_name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(10.5)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
