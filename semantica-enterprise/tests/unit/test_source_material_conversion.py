"""Conversion contracts use synthetic documents, not customer attachments."""
from pathlib import Path
import importlib.util

MODULE = Path(__file__).resolve().parents[2] / "scripts/verification/convert_source_materials.py"
spec = importlib.util.spec_from_file_location("source_material_conversion", MODULE)
conversion = importlib.util.module_from_spec(spec)
spec.loader.exec_module(conversion)


def test_docx_merged_cells_not_duplicated_and_identical_distinct_cells_kept(tmp_path):
    from docx import Document
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("建设规模", 1)
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并标题"
    table.cell(0, 2).text = "备注"
    table.cell(1, 0).text = "10"
    table.cell(1, 1).text = "10"
    table.cell(1, 2).text = "数值"
    doc.save(source)
    before = source.read_bytes()
    content, audit = conversion.convert_docx(source)
    assert content.count("合并标题") == 1
    assert "| 10 | 10 | 数值 |" in content
    assert "## 建设规模" in content and audit["tables"] == 1
    assert source.read_bytes() == before
