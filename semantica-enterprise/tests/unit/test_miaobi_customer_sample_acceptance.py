from pathlib import Path

from openpyxl import Workbook

from scripts.miaobi.acceptance_customer_sample import (
    EXPECTED_COMPUTATIONS,
    REQUIRED_FACTS,
    content_quality_metrics,
    extract_workbook_facts,
    markdown_to_plate,
    renumber_citations,
)


def test_extract_workbook_facts_uses_named_sheets_and_exact_rows(tmp_path: Path):
    path = tmp_path / "facts.xlsx"
    workbook = Workbook()
    index = workbook.active
    index.title = "00_数据引用索引"
    index.append(["数据ID", "数据项", "值", "单位"])
    for row in [
        ["D01", "发生时间", "12月18日23时59分", None],
        ["D02", "震中", "甘肃临夏州积石山县柳沟乡附近", None],
        ["D04", "震级", 6.2, "级"],
        ["D05", "震源深度", 10, "km"],
        ["D07", "最高烈度", "VIII", "度"],
    ]:
        index.append(row)
    grading = workbook.create_sheet("R01_灾害分级与响应推演")
    grading.append(["说明"])
    grading.append(["推演步骤", "输入", "计算", "结论"])
    grading.append(["判据区域成立性", "人口与面积", "278089÷909.97 = 305.6 > 200 人/km²", "成立"])
    resources = workbook.create_sheet("R03_资源需求缺口总表")
    resources.append(["说明"])
    resources.append(["资源", "单位", "公式", "过程", "目标", "县域", "可调用", "缺口"])
    resources.append(["搜救人员", "人", "", "", 500, 320, 320, 180])
    resources.append(["创伤床位", "张", "", "", 330, 110, 250, "县域220 / 全域80"])
    resources.append(["帐篷", "顶", "", "", 7000, 5200, 5200, 1800])
    workbook.save(path)

    facts = extract_workbook_facts(path)

    assert set(REQUIRED_FACTS).issubset(facts)
    assert facts["magnitude"]["value"] == {"number": 6.2}
    assert facts["population_density"]["value"] == {"number": 305.6}
    assert facts["callable_trauma_beds"]["value"] == {"number": 250}
    assert facts["magnitude"]["sheet"] == "00_数据引用索引"
    assert facts["rescue_required"]["sheet"] == "R03_资源需求缺口总表"
    assert EXPECTED_COMPUTATIONS == {
        "rescue_gap": 180,
        "county_bed_gap": 220,
        "all_area_bed_gap": 80,
        "tents_gap": 1800,
    }


def test_renumber_and_plate_materialization_create_native_inline_references():
    messages = [
        {
            "content": "## 一、总则\n\n依据预案执行[1]。",
            "citations": [
                {
                    "citation_number": 1,
                    "query_run_id": "query-1",
                    "chunk_id": "chunk-1",
                    "snapshot": {"title": "预案", "document_id": "doc-1", "version_id": "v1"},
                }
            ],
        },
        {
            "content": "## 二、震情灾情\n\n继续使用同一依据[1]，并增加数据表[2]。",
            "citations": [
                {
                    "citation_number": 1,
                    "query_run_id": "query-1",
                    "chunk_id": "chunk-1",
                    "snapshot": {"title": "预案", "document_id": "doc-1", "version_id": "v1"},
                },
                {
                    "citation_number": 2,
                    "query_run_id": "query-2",
                    "chunk_id": "chunk-2",
                    "snapshot": {"title": "数据表", "document_id": "doc-2", "version_id": "v2"},
                },
            ],
        },
    ]

    normalized, evidence = renumber_citations(messages)
    nodes, bindings = markdown_to_plate(normalized, evidence)

    assert len(evidence) == 2
    assert all("<" not in str(node) for node in nodes)
    assert any(node.get("type") == "h2" and node["children"][0]["text"] == "一、总则" for node in nodes)
    assert len(bindings) == 3
    assert all(item["reference"]["type"] == "knowledge_citation" for item in bindings)
    assert all(item["reference"]["children"] == [{"text": ""}] for item in bindings)


def test_customer_acceptance_rejects_work_notes_duplicates_letter_and_missing_trusted_blocks(tmp_path: Path):
    from docx import Document

    docx_path = tmp_path / "unsafe.docx"
    exported = Document()
    exported.add_paragraph("一、总则")
    exported.add_paragraph("Let me compose the section. 正文没有可见引用。")
    exported.save(docx_path)
    nodes = [
        {"type": "h2", "children": [{"text": "一、总则"}]},
        {"type": "p", "children": [{"text": "Let me compose the section. 图谱查询返回为空。"}]},
        {"type": "h2", "children": [{"text": "一、总则"}]},
        {"type": "p", "children": [{"text": "这是被重复写入的同一段正式方案内容，不能通过生产质量门禁。"}]},
        {"type": "p", "children": [{"text": "这是被重复写入的同一段正式方案内容，不能通过生产质量门禁。"}]},
    ]

    metrics = content_quality_metrics(
        nodes,
        [{"format": "docx", "path": str(docx_path)}],
    )

    assert metrics["private_work_matches"] > 0
    assert metrics["platform_process_matches"] > 0
    assert metrics["chapter_occurrences"]["一、总则"] == 2
    assert metrics["duplicate_sentence_count"] == 1
    assert metrics["docx"]["is_a4"] is False
    assert metrics["docx"]["visible_inline_citations"] == 0
    assert metrics["computed_metric_nodes"] == 0
    assert metrics["inference_conclusion_nodes"] == 0
