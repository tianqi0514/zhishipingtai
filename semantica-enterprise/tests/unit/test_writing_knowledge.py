from copy import deepcopy
import json

import pytest

from packages.platform.writing_knowledge import build_chapter_packet
from packages.platform.writing import affected_dependency_ids, validate_plate_content
from packages.platform.writing_flow import assemble_report_content, validate_and_parse_agent_report
from packages.platform.writing_configuration import compile_business_scenario, business_scenario_from_contract
from apps.api.writing_schemas import ScenarioSectionSetting


def sample_facts():
    def f(i, subject, name, predicate, obj, obj_name, typ, space="space"):
        return {"id": i, "space_id": space, "subject_entity_id": subject,
                "subject_name": name, "subject_type": typ, "predicate": predicate,
                "object_entity_id": obj, "object_name": obj_name,
                "object_value": None, "confidence": 1,
                "source_chunk_id": f"chunk-{i}",
                "source": {"chunk_id": f"chunk-{i}", "title": f"材料{i}",
                           "document_id": f"doc-{i}", "version_id": f"version-{i}", "page_number": 1}}
    return [f("1", "shelter", "安置点A", "依赖", "water", "供水站C", "安置点"),
            f("2", "water", "供水站C", "发生", "outage", "供水中断", "供水设施"),
            f("3", "dept", "保障组D", "负责", "shelter", "安置点A", "责任单位")]


RULE = {"id": "risk-rule", "version_id": "rule-v1", "name": "供水风险传导", "confidence": 1,
        "definition": {"conditions": [{"predicate": "依赖", "subject": "X", "object": "Y"},
                                      {"predicate": "发生", "subject": "Y", "object": "Z"}],
                       "conclusion": {"predicate": "受到影响", "subject": "X", "object": "Z"}}}
CHAPTER = {"key": "water", "title": "安置点饮水保障", "knowledge": {"entity_types": ["安置点"], "predicates": ["依赖", "发生", "负责"], "rule_version_ids": ["rule-v1"]}}


def packet(facts=None):
    return build_chapter_packet(chapters=[CHAPTER], facts=facts if facts is not None else sample_facts(), requirements={"water": ({}, [RULE])})


def test_real_semantica_derives_missing_cross_document_risk_with_sources():
    result = packet()
    risk = [r for r in result["references"].values() if r["kind"] == "inference"]
    assert len(risk) == 1
    assert risk[0]["text"] == "安置点A — 受到影响 → 供水中断"
    assert {s["chunk_id"] for s in risk[0]["sources"]} == {"chunk-1", "chunk-2"}
    assert risk[0]["proof"]["engine"] == "semantica.reasoning.DatalogReasoner"
    assert result["sections"][0]["engine_metrics"][0]["rules"] == 1
    assert any("保障组D" in r["text"] for r in result["references"].values())


def test_restoration_recomputes_no_negative_fact_and_keeps_old_packet():
    before = packet()
    facts = sample_facts()
    facts[1]["predicate"] = "已恢复"
    after = packet(facts)
    assert any(r["kind"] == "inference" for r in before["references"].values())
    assert not any(r["kind"] == "inference" for r in after["references"].values())
    assert "缺少“发生”关系的来源依据" in after["sections"][0]["warnings"]


def test_rule_cannot_join_facts_across_spaces():
    facts = sample_facts()
    facts[1]["space_id"] = "other"
    assert packet(facts)["sections"][0]["conclusion_count"] == 0


def test_missing_object_does_not_claim_ready():
    facts = sample_facts()
    facts[0]["subject_type"] = "其他对象"
    assert "已选材料中没有匹配的业务对象" in packet(facts)["sections"][0]["warnings"]


def test_same_inputs_generate_stable_reference_identity():
    assert set(packet()["references"]) == set(packet(list(reversed(sample_facts())))["references"])


def test_ontology_alias_normalization_does_not_modify_input():
    facts = sample_facts()
    facts[0]["predicate"] = "depend_on"
    original = deepcopy(facts)
    result = build_chapter_packet(chapters=[CHAPTER], facts=facts, requirements={"water": ({"depend_on": "依赖"}, [RULE])})
    assert result["sections"][0]["conclusion_count"] == 1
    assert facts == original


def test_paragraph_dependencies_become_stale_without_modifying_prose():
    content = [{"id": "paragraph", "type": "p", "children": [{"text": "待增派180名搜救人员。"}]}]
    binding = {"block_id": "paragraph", "metadata": {"input_fact_ids": ["available"], "computation_run_ids": ["calc"]}, "freshness_status": "stale"}
    impact = affected_dependency_ids({"available"}, [{"id": "calc", "input_fact_ids": ["available"]}], [binding])
    assert impact["block_ids"] == ["paragraph"]
    assert validate_plate_content(content, {"paragraph": binding})[0]["code"] == "stale_paragraph"
    assert "180" in content[0]["children"][0]["text"]


def assemble(text, **extra):
    return assemble_report_content(run_id="run", title="保障方案", section_plan=[CHAPTER],
        agent_sections=[{"section_key": "water", "content_nodes": [{"type": "p", "text": text, **extra}]}],
        citations=[], computations=[], inference_facts=[], selected_plan=None, knowledge_packet={**packet(), "run_id": "knowledge-run"}, input_facts=[])


def test_semantic_reference_appears_in_clean_prose_with_proof_in_binding():
    ref = next(r for r, v in packet()["references"].items() if v["kind"] == "inference")
    content, bindings = assemble(f"安置点存在供水中断影响，需协调替代供水。[{ref}]")
    paragraph = next(b for b in bindings if b["block_type"] == "p")
    assert paragraph["metadata"]["knowledge_refs"] == [ref]
    assert paragraph["verification_status"] == "unverified"
    assert paragraph["metadata"]["knowledge_evidence"][0]["sources"]
    assert any(b["block_type"] == "knowledge_citation" for b in bindings)
    assert not any(i["code"] in ("invalid_source_type", "missing_binding") for i in validate_plate_content(content, {b["block_id"]: b for b in bindings}))


@pytest.mark.parametrize("text,extra", [("不能使用不存在的依据[K000000000000]", {}), ("不存在的文档[9]", {}), ("自行指定输入", {"input_refs": ["unknown"]}), ("自行指定计算", {"metric_refs": ["unknown"]})])
def test_fabricated_references_fail_closed(text, extra):
    with pytest.raises(ValueError):
        assemble(text, **extra)


def test_business_config_round_trip_keeps_real_chapter_requirements():
    payload = {"inputs": [{"key": "population", "label": "人数", "data_type": "number"}],
               "sections": [{**CHAPTER, "purpose": "说明饮水保障", "knowledge": {**CHAPTER["knowledge"], "ontology_version_id": "ontology-v1"}}]}
    compiled = compile_business_scenario(payload)
    assert business_scenario_from_contract(compiled)["sections"][0]["knowledge"] == payload["sections"][0]["knowledge"]


def test_section_unknown_knowledge_fields_rejected():
    with pytest.raises(ValueError):
        ScenarioSectionSetting.model_validate({"key": "water", "title": "供水", "purpose": "供水保障", "knowledge": {"raw_sql": "select 1"}})


def test_configured_chapter_cannot_silently_ignore_its_relationships():
    with pytest.raises(ValueError, match="未引用"):
        assemble("一般来说，应做好保障工作。")


def test_malformed_semantic_citation_is_rejected():
    with pytest.raises(ValueError, match="没有实际依据"):
        assemble("不能捏造依据[Knot_real]")


@pytest.mark.parametrize("status,expected", [("failed", "agent_failed"), ("cancelled", "cancelled"), ("completed", "agent_running")])
def test_failed_agent_projection_does_not_leave_generation_running(status, expected):
    from types import SimpleNamespace
    from apps.api.writing import _reconcile_generation_failure
    assistant = SimpleNamespace(status=status, error_code="NETWORK", error_message="网络不可用")
    run = SimpleNamespace(status="agent_running", assistant_message_id="message")
    db = SimpleNamespace(get=lambda *args: assistant)
    _reconcile_generation_failure(db, run)
    assert run.status == expected


def test_explicit_empty_dependencies_do_not_mark_unrelated_paragraphs():
    content, bindings = assemble_report_content(run_id="run", title="保障", section_plan=[{"key": "s", "title": "人员", "required_inputs": ["available"]}],
        agent_sections=[{"section_key": "s", "content_nodes": [{"type": "p", "text": "交通情况有待核实。", "input_refs": [], "metric_refs": []}]}],
        citations=[], computations=[], inference_facts=[], selected_plan=None, input_facts=[{"id": "f", "fact_key": "available", "value": {"number": 400}}])
    assert bindings[0]["metadata"]["input_fact_ids"] == []


def test_formal_prose_rejects_engine_execution_labels():
    payload = {"sections": [{"section_key": "water", "title": CHAPTER["title"],
        "content_nodes": [{"type": "p", "text": "规则推演结论为：安置点受到影响供水中断。"}]}]}
    with pytest.raises(ValueError, match="平台执行说明"):
        validate_and_parse_agent_report(json.dumps(payload, ensure_ascii=False), [CHAPTER])


def test_publication_rejects_stale_narrative(monkeypatch):
    from types import SimpleNamespace
    from fastapi import HTTPException
    from apps.api import writing as api
    from apps.api.writing_schemas import WritingDocumentVersionCreate
    monkeypatch.setattr(api, "_document", lambda *a: (SimpleNamespace(id="doc"), SimpleNamespace(id="project")))
    monkeypatch.setattr(api, "_generated_report_quality", lambda *a, **kw: None)
    monkeypatch.setattr(api, "validate_plate_content", lambda *a: [{"code": "stale_paragraph", "block_id": "p"}])
    db = SimpleNamespace(scalars=lambda *a: [], scalar=lambda *a: 0)
    with pytest.raises(HTTPException) as error:
        api.create_document_version("doc", WritingDocumentVersionCreate(content=[], publish=True), SimpleNamespace(), db)
    assert error.value.status_code == 409


def test_prose_revision_rejects_execution_notes_and_unrequested_citations():
    from packages.platform.writing_flow import validate_agent_edit

    original = "保障组先核实库存和运输条件，再协调临时送水；车辆数量和到达时间须现场确认。"
    assert validate_agent_edit("shorten", original, "保障组核实库存与运输后协调送水，车辆及到达时间待现场确认。")
    with pytest.raises(ValueError, match="工具执行说明"):
        validate_agent_edit("shorten", original, "已调用写作工具。车辆待确认。")
    with pytest.raises(ValueError, match="原文没有的引用"):
        validate_agent_edit("shorten", original, "现场核实后送水。[2]")
    with pytest.raises(ValueError, match="编写指令"):
        validate_agent_edit("shorten", original, "车辆待确认，不得凭空生成。")
    assert validate_agent_edit("shorten", original + "[2]", "保障组协调送水，车辆及时间待现场确认。[2]")


def test_export_preserves_all_inference_premises_without_internal_ids(tmp_path):
    from docx import Document
    from packages.platform.writing_export import build_evidence_docx
    references = packet()["references"]
    inference = next(r for r in references.values() if r["kind"] == "inference")
    binding = {"block_type": "knowledge_citation", "metadata_json": {
        "citation_number": 7, "knowledge_evidence": [inference]}}
    target = tmp_path / "evidence.docx"
    build_evidence_docx(target, title="安置保障", facts=[], computations=[], plans=[], audit_summary={}, bindings=[binding, binding])
    doc = Document(target)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.count("[7] 规则结论") == 1
    assert "安置点A — 依赖 → 供水站C" in text
    assert "供水站C — 发生 → 供水中断" in text
    assert "材料1" in text and "材料2" in text and "第 1 页" in text
    assert "rule-v1" not in text and "chunk-1" not in text
    assert any(r.cells[0].text == "正文规则结论" and r.cells[1].text == "1" for r in doc.tables[0].rows)


def test_formal_export_does_not_invent_subtitle_or_platform_opening(tmp_path):
    from docx import Document
    from packages.platform.writing_export import build_docx
    target = tmp_path / "formal.docx"
    build_docx(target, title="供水保障报告", content=[{"type": "p", "children": [{"text": "先核实库存。"}]}], audit_summary={})
    assert [p.text for p in Document(target).paragraphs if p.text] == ["供水保障报告", "先核实库存。"]


def test_export_uses_current_block_identity_not_reused_historical_citation_number(tmp_path):
    from docx import Document
    from packages.platform.writing_export import bindings_for_content, build_evidence_docx
    content = [{"id": "p", "children": [{"id": "current", "type": "knowledge_citation", "children": [{"text": ""}]}]}]
    bindings = {"old": {"block_type": "knowledge_citation", "metadata_json": {"citation_number": 10, "source_title": "历史错误材料"}},
                "current": {"block_type": "knowledge_citation", "metadata_json": {"citation_number": 10, "source_title": "当前巡查简报"}},
                "old-metric": {"block_type": "computed_metric", "computation_run_id": "old-run"}}
    selected = bindings_for_content(content, bindings)
    assert set(selected) == {"current"}
    target = tmp_path / "evidence.docx"
    build_evidence_docx(target, title="报告", facts=[], computations=[], plans=[], audit_summary={}, bindings=list(selected.values()))
    text = "\n".join(p.text for p in Document(target).paragraphs)
    assert "[10] 当前巡查简报" in text
    assert "历史错误材料" not in text
