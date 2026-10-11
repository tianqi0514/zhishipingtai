from docx import Document
from docx.oxml.ns import qn

from packages.platform.writing_export import _add_table


def test_export_table_repeats_only_first_row_header():
    document = Document()
    node = {
        "type": "table",
        "children": [
            {"type": "tr", "children": [
                {"type": "th", "children": [{"text": "费用分项"}]},
                {"type": "th", "children": [{"text": "金额（万元）"}]},
            ]},
            {"type": "tr", "children": [
                {"type": "td", "children": [{"text": "项目总投资"}]},
                {"type": "td", "children": [{"text": "12800"}]},
            ]},
        ],
    }
    _add_table(document, node)
    table = document.tables[0]
    assert table.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
    assert not table.rows[1]._tr.xpath("./w:trPr/w:tblHeader")
    assert table.cell(1, 1).text == "12800"


def test_empty_export_table_does_not_create_a_header():
    document = Document()
    _add_table(document, {"type": "table", "children": []})
    assert not document.tables
