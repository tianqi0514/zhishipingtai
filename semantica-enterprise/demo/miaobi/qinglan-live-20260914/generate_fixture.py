#!/usr/bin/env python3
"""Rebuild only this isolated fictional exercise fixture with bundled Python.

Authoring source of truth lives in this file. Run from any working directory.
Final inputs are under 01-上传材料; change-only material is outside that folder.
"""

from __future__ import annotations

import json
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

BASE = Path(__file__).resolve().parent
INITIAL = BASE / "01-上传材料"
CHANGE = BASE / "02-变更参考"
DISCLAIMER = "演示材料  所有地名和业务记录均属虚构，不代表真实灾情，不构成调度命令。"
BODY_FONT = "Songti SC"
HEAD_FONT = "Heiti SC"

FACILITY = {
    "title": "青岚县地震演练安置点设施台账",
    "meta": "编制单位：县安置设施登记小组  |  台账编号：QL-AZ-0914-01",
    "version": "版本 V1.0  |  记录时点：2026-09-14 08:20  |  适用范围：本次演练",
    "intro": "本台账记录本次地震演练中两个安置点的登记规模和生活用水设施关系，供保障人员核对基础信息。登记小组以演练设施底册和当班登记为依据建立本版记录；本台账不替代设施实时巡查，不对供水连续性、水质或临时运输能力作出保证。",
    "sections": [
        ("一 登记范围与数据口径", [
            "台账中的人数为08:20登记在册的安置人数，单位为人，分别对应各安置点，不代表搜救作业任务量、救援编组或需增派人员数量。人员流动后应由登记小组按新时点更新，不得把不同时间的登记数混用。设施关系记录的是本版已登记的供水来源，并非对区域管网完整拓扑的认定。",
        ]),
        ("二 安置点设施登记", []),
        ("三 使用边界与待核事项", [
            "云桥体育馆安置点的生活用水依赖柳溪供水站，底册登记的站点供水接入关系作为后续核查依据。现有材料没有给出该安置点独立备用水源的可用能力，不能据此假定备用保障已经落实。点内存水量、分发设施的可用状态以及可持续保障时长，应由现场人员另行核实并注明核查时间。",
            "城北中学安置点由城北备用水井供水。当前设施底册和当班登记未发现其依赖柳溪供水站的记录；这一表述仅说明本版掌握的设施关系，不等于证明城北中学安置点没有供水风险。备用水井的运行状态、水质检测及其他潜在关联仍须依照巡查记录核实，未取得确认前不得写成正常或安全。",
            "后续如发现接入关系变更，登记小组应保留原记录，以新版本注明变更依据和生效时点。各业务组引用本台账时应保留设施名称和记录时点；涉及运行异常的判断，应同时查阅相应设施的最新巡查材料。本版未记载任何供水车辆出发、到达或物资签收信息。",
        ]),
    ],
    "table": {
        "headers": ["安置点", "登记人数", "生活用水来源", "登记依据"],
        "rows": [
            ["云桥体育馆安置点", "600人", "柳溪供水站", "演练设施底册及当班登记"],
            ["城北中学安置点", "420人", "城北备用水井", "演练设施底册及当班登记"],
        ],
    },
}

INCIDENT = {
    "title": "柳溪供水站巡查简报",
    "meta": "报送单位：供水巡查组  |  简报编号：QL-GS-0914-01",
    "version": "版本 V1.0  |  确认时点：2026-09-14 09:00  |  适用范围：本次演练",
    "intro": "供水巡查组确认，柳溪供水站在本次地震演练情境中于2026年9月14日08:40发生供水中断，09:00完成状态复核。截至本简报确认时点，该站供水中断状态成立，恢复时间尚未确认。请接收单位使用这一时点明确的设施状态，不将后续计划表述为已完成事项。",
    "sections": [
        ("一 巡查记录与确认范围", [
            "本次记录对象为柳溪供水站。08:40为中断发生时间，09:00为巡查组复核确认时间，两者含义不同。巡查组按演练记录核对供水中断情况并形成此报，记录覆盖站点运行状态，不包含用户端逐户查验，也未完成对下游用水对象的逐项核对。设施名称应按本简报原名引用，避免因简称不同漏记状态。",
            "本简报不据中断现象认定故障部位、损坏程度或确切成因。相关技术排查和恢复条件仍需核实，未取得进一步记录前，应将这些事项保留为待核状态，不得自行补写管线破裂、设备报废或已恢复运行等结论。",
        ]),
        ("二 尚未确认的保障条件", [
            "截至09:00，供水恢复的预计时间和验证结果均未确认，不能承诺某一时刻恢复供水。水质情况尚未取得确认材料，不能因供水停止而直接判断污染，也不能将既往状态视为当前水质合格的依据。恢复供水与水质可用应分别记录和核查。",
            "替代供水及运输条件尚未确认。本简报没有记载可用送水车辆、装载量、行车路线、发车时间或到达回执，不能据此写出运输已安排或水已送达。替代水源的可用量及取水条件也需另行核实，拟议措施只有取得相应记录后才能更新为实际进展。",
        ]),
        ("三 续报要求", [
            "供水巡查组负责柳溪供水站的续报，后续材料应分别说明核实时间、设施状态变化及仍未确认的事项，并保留本次中断的起始时间。接收单位如掌握新的现场记录，应先核对对象和时点，再据以更新业务材料；本简报本身不发布调度决定。",
        ]),
    ],
}

GUIDE = {
    "title": "青岚县地震演练生活保障职责与工作指引",
    "meta": "编制单位：演练指挥协调组  |  指引编号：QL-BZ-0914-01",
    "version": "版本 V1.0  |  适用起点：2026-09-14 08:00  |  适用范围：本次演练",
    "intro": "本指引用于明确演练中的生活用水保障职责及记录要求。保障工作应先核实现场条件，再提出可执行的临时供水安排；设施运行状态与安置点保障状态分别记录。本指引规定处理流程，不说明某一设施已发生异常，也不证明某项保障措施已经实施。",
    "sections": [
        ("一 职责分工", [
            "县生活保障组负责云桥体育馆安置点的生活保障协调，汇总点内用水、库存和补充需求，组织核对保障条件，并向演练指挥协调组提交情况及拟办事项。点位名称应使用全称，保障对象不得因材料简称相近而被替换。",
            "供水巡查组负责柳溪供水站的运行巡查、状态核实和变化续报，说明设施状态的记录时点及恢复条件。巡查组提供设施侧事实，不以巡查记录代替安置点现场盘点，也不代替县生活保障组决定具体供水安排。",
        ]),
        ("二 临时供水前的核实事项", [
            "县生活保障组在提出临时供水安排前，应核实点内实际库存及可取用状态，区分已入库、在途和仅列入计划的水量。库存持续时长如需估算，应同时注明测算依据，不能仅按登记人数推定库存已经不足或能够维持某一时段。",
            "运输核实应覆盖实际可用运力、取水点条件、通行情况和接收能力。水质核实应取得对应水源及使用环节的有效确认，不以供水恢复、运输获批或外观正常代替水质结论。库存、运输和水质条件未核清的部分，应分别列入待办事项并说明负责核实的业务组。",
        ]),
        ("三 安排与完成状态的记录", [
            "县生活保障组在核实库存、运输和水质后，结合实际缺项提出临时供水安排，并按演练协调程序落实。安排记录应区分拟议、确认、执行和签收状态；只有取得相应凭据，才能写明车辆已出发、水已到点或已完成分发。指引没有指定车辆、运量和到达时刻。",
        ]),
        ("四 信息更新与待核管理", [
            "保障记录应保留资料来源和有效时点。设施状态变化由供水巡查组续报，点内条件变化由县生活保障组汇总。遇到资料缺失或口径不一致，应先提出核实事项；报告可说明当前约束和下一步动作，但不得把待核条件写成既成事实。本指引不授权真实调度。",
        ]),
    ],
}

RESCUE_MD = """# 青岚县地震演练搜救力量需求及到位清单

演示材料：所有地名和业务记录均属虚构，不代表真实灾情，不构成调度命令。

编制单位：演练指挥协调组力量登记席  
清单编号：QL-SJ-0914-01；版本：V1.0  
需求口径：本次演练县域搜救专项任务；到位时点：2026-09-14 09:00

## 一 需求依据

本清单将县域搜救专项任务需求与09:00实际可用力量分别列示，供力量核对和报告编写使用。需求来自演练指挥协调组任务单QL-RW-0914-01，任务单核定总需求为500人，按县救援力量、专业力量和机动力量分类列明。该需求是本次演练任务设定，不是根据安置人数测算，也不得用安置点登记人数替换。

本清单只涉及搜救专项人员，不包括生活供水、医疗、物资分发等其他业务的人员配置。需求总量及分类口径在取得任务单变更前保持不变；不能将力量到位变化理解为任务需求自动减少。

## 二 需求与可用记录

| 力量类别 | 任务单需求 | 09:00实际可用 |
| --- | ---: | ---: |
| 县救援力量 | 240人 | 160人 |
| 专业力量 | 160人 | 100人 |
| 机动力量 | 100人 | 60人 |
| 合计 | 500人 | 320人 |

“实际可用”仅指截至09:00已经到位、完成登记并纳入本次搜救专项可用记录的人员。同一人员不得在两个力量类别中重复计数。表中需求与可用均以人为单位，指向同一县域搜救专项范围，引用时应保留时点，避免将人数与队伍数混淆。

## 三 后备人员与待核事项

另有后备人员60人，截至09:00尚未到场，未计入表中320人。这60人属于待到位记录，不是已经可用人员，也不是在320人之外已经到场的另一批力量。后备编组和联络记录仅用于跟踪准备状态，不能作为到场证明。

力量登记席后续应依据实际到位确认更新清单，并核对重复登记、离场和任务转派情况。尚未获得到场确认的人员继续保留在待到位记录中，不得因预计到达而提前加入可用总数。本版未对后续到达时间作出承诺，也未记载任何新增到位事项。

## 四 引用和更新要求

使用本清单形成分析时，应先核对需求和可用力量是否属于相同任务范围、相同单位及指定时点，再计算需要报告的数值。更新版应说明变化发生时间、所依据的到位记录以及是否影响任务单需求，保留本版作为09:00历史记录。本清单提供统计事实，不代替力量调度决定。
"""

UPDATE_MD = """# 青岚县地震演练搜救人员到位更新

演示材料：所有地名和业务记录均属虚构，不代表真实灾情，不构成调度命令。

编制单位：演练指挥协调组力量登记席  
更新编号：QL-SJ-0914-02；版本：V2.0  
确认时点：2026-09-14 11:00；关联记录：QL-SJ-0914-01

## 一 本次更新范围

截至2026年9月14日11:00，力量登记席确认本次县域搜救专项新增到位80人，已完成登记并纳入实际可用记录。本次更新只改变搜救人员的到位及可用数量，不改变搜救任务范围，也不表示其他业务保障条件发生变化。09:00版本仍用于查询对应历史时点，不能把两个版本的可用总数相加。

## 二 新增到位记录

| 新增人员来源 | 11:00新增到位 | 与09:00记录的关系 |
| --- | ---: | --- |
| 后备人员 | 60人 | 09:00尚未到场且未计入320人的原后备人员 |
| 新增增援人员 | 20人 | 本次更新确认的新增到位人员 |
| 新增合计 | 80人 | 本次新增，未重复计入09:00可用人数 |

上述后备60人已由待到位记录转为实际可用记录，不得继续同时保留在待到位合计中，也不得在新增80人之外再次加计。新增增援20人与原后备60人为不同记录来源，本次统计已经核对两者不重叠。

## 三 更新后的统计口径

本次更新以09:00实际可用320人为基数，加上11:00新增到位80人，11:00实际可用总数为400人。本次更新口径中，原320人继续可用，没有新增退出或重复扣减事项。任务单QL-RW-0914-01核定的县域搜救专项总需求仍为500人，未发生调整。

本次新增人员仅按来源登记，尚未提供其分入县救援力量、专业力量和机动力量的分类明细。报告可以使用已确认的可用总数，不应自行补出更新后的各类别人数。需求分类仍以原任务单为准，人员来源与任务类别不是同一统计维度。

## 四 续报与引用要求

力量登记席后续应按实际到位、离场及转派记录继续更新，遇到分类明细缺失应列为待核事项。引用本更新时，应注明11:00时点和需求未变的前提；本更新未记载任何设施状态、供水恢复、运输安排或安置点人数变化，不应据其推断相关业务进展。此更新供演练信息校核使用，不代替真实调度命令。
"""


def east_asian_font(style, name):
    style.font.name = name
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for field in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{field}"), name)


def add_field(p, instruction):
    run = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruction)
    run._r.addnext(fld)


def make_docx(data, output):
    doc = Document()
    # The bundled base document may contain a themed Title bottom border.
    # Remove inherited paragraph borders, including the blue title rule.
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.63), Inches(0.60)
    sec.left_margin, sec.right_margin = Inches(0.72), Inches(0.72)
    sec.footer_distance = Inches(0.26)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"):
        east_asian_font(doc.styles[name], BODY_FONT if name in ("Normal", "Footer") else HEAD_FONT)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.widow_control = True
    title = doc.styles["Title"]
    title.font.size = Pt(19)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.line_spacing = 1.10
    heading = doc.styles["Heading 1"]
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.paragraph_format.space_before = Pt(9)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(data["title"], "Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in (data["meta"], data["version"]):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(9)
    p = doc.add_paragraph(DISCLAIMER)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.size = Pt(9)
    doc.add_paragraph(data["intro"])
    for heading_text, paragraphs in data["sections"]:
        doc.add_paragraph(heading_text, "Heading 1")
        if heading_text == "二 安置点设施登记":
            add_table(doc, data["table"])
        for text in paragraphs:
            p = doc.add_paragraph(text)
            p.paragraph_format.first_line_indent = Pt(22)
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run("演示材料  |  ")
    add_field(foot, "PAGE")
    for run in foot.runs:
        run.font.size = Pt(8)
    doc.core_properties.author = "青岚县地震演练材料组（虚构）"
    doc.core_properties.title = data["title"]
    doc.core_properties.subject = "虚构演练材料，不代表真实灾情，不构成调度命令"
    doc.core_properties.comments = ""
    doc.save(output)


def add_table(doc, data):
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    widths = (1.92, 0.80, 1.63, 2.71)
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for i, text in enumerate(data["headers"]):
        table.rows[0].cells[i].text = text
    for row in data["rows"]:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    for row_idx, row in enumerate(table.rows):
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{edge}")
                for k, v in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    el.set(qn(f"w:{k}"), v)
                borders.append(el)
            tcPr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:w"), "90")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            tcPr.append(margins)
            if row_idx == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8ECF0")
                tcPr.append(shade)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 1 or row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = row_idx == 0
        if row_idx == 0:
            header = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(header)


def make_pdf(data, output):
    pdfmetrics.registerFont(TTFont("QinglanChinese", "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"))
    style = dict(fontName="QinglanChinese", textColor=colors.black, wordWrap="CJK")
    title = ParagraphStyle("Title", fontSize=18, leading=24, alignment=TA_CENTER, spaceAfter=9, **style)
    meta = ParagraphStyle("Meta", fontSize=8.5, leading=13, alignment=TA_CENTER, spaceAfter=3, **style)
    note = ParagraphStyle("Note", fontSize=9, leading=14, spaceAfter=10, **style)
    body = ParagraphStyle("Body", fontSize=11, leading=17, spaceAfter=7, firstLineIndent=22, alignment=TA_JUSTIFY, **style)
    intro = ParagraphStyle("Intro", parent=body, firstLineIndent=0)
    heading = ParagraphStyle("Heading", fontSize=12, leading=18, spaceBefore=8, spaceAfter=5, keepWithNext=True, **style)
    story = [Paragraph(escape(data["title"]), title), Paragraph(escape(data["meta"]), meta),
             Paragraph(escape(data["version"]), meta), Spacer(1, 5), Paragraph(escape(DISCLAIMER), note),
             Paragraph(escape(data["intro"]), intro)]
    for label, paragraphs in data["sections"]:
        story.append(Paragraph(escape(label), heading))
        story.extend(Paragraph(escape(text), body) for text in paragraphs)
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("QinglanChinese", 8)
        canvas.drawCentredString(letter[0] / 2, 23, f"演示材料  |  {document.page}")
        canvas.restoreState()
    pdf = SimpleDocTemplate(str(output), pagesize=letter, rightMargin=52, leftMargin=52,
                            topMargin=42, bottomMargin=42, title=data["title"],
                            author="青岚县地震演练材料组（虚构）", subject=DISCLAIMER)
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)


def content_text(data):
    parts = [data["title"], data["meta"], data["version"], DISCLAIMER, data["intro"]]
    for title, paragraphs in data["sections"]:
        parts.extend([title, *paragraphs])
    if "table" in data:
        parts.extend(data["table"]["headers"])
        parts.extend(value for row in data["table"]["rows"] for value in row)
    return "\n".join(parts)


def main():
    INITIAL.mkdir(parents=True, exist_ok=True)
    CHANGE.mkdir(parents=True, exist_ok=True)
    make_docx(FACILITY, INITIAL / "01-安置点设施台账.docx")
    make_pdf(INCIDENT, INITIAL / "02-柳溪供水站巡查简报.pdf")
    make_docx(GUIDE, INITIAL / "03-生活保障职责与工作指引.docx")
    (INITIAL / "04-搜救力量需求及到位清单.md").write_text(RESCUE_MD, encoding="utf-8")
    (CHANGE / "05-人员到位更新.md").write_text(UPDATE_MD, encoding="utf-8")
    truth = {
        "fixture_id": "qinglan-live-20260914",
        "synthetic": True,
        "warning": DISCLAIMER,
        "scope": "仅本目录独立演示，不含真实客户资料、真实灾情或真实个人信息",
        "timezone": "Asia/Shanghai",
        "initial_upload_dir": "01-上传材料",
        "change_reference_dir": "02-变更参考",
        "not_for_upload": ["source_truth.json", "generate_fixture.py", "README.md", "qa"],
        "facts": {
            "facility_0820": {
                "source": "01-上传材料/01-安置点设施台账.docx",
                "as_of": "2026-09-14T08:20:00+08:00",
                "云桥体育馆安置点": {"registered_population": 600, "depends_on": "柳溪供水站", "independent_backup_capacity": "unknown"},
                "城北中学安置点": {"registered_population": 420, "supplied_by": "城北备用水井", "dependency_on_柳溪供水站": "not_found_in_current_records", "risk_free": "not_established"},
                "population_is_rescue_demand_basis": False,
            },
            "water_incident": {
                "source": "01-上传材料/02-柳溪供水站巡查简报.pdf",
                "facility": "柳溪供水站", "status": "供水中断",
                "occurred_at": "2026-09-14T08:40:00+08:00", "confirmed_at": "2026-09-14T09:00:00+08:00",
                "restoration_time": "unknown", "water_quality": "unknown", "alternate_transport": "unknown",
                "cause": "unknown", "downstream_shelter_names_in_source": [],
            },
            "responsibilities": {
                "source": "01-上传材料/03-生活保障职责与工作指引.docx",
                "县生活保障组": "云桥体育馆安置点生活保障协调",
                "供水巡查组": "柳溪供水站运行巡查 状态核实 变化续报",
                "temporary_supply_prerequisites": ["核实库存", "核实运输", "核实水质"],
                "transport_dispatched_or_arrived": "not_recorded",
            },
            "rescue_0900": {
                "source": "01-上传材料/04-搜救力量需求及到位清单.md",
                "as_of": "2026-09-14T09:00:00+08:00", "unit": "人", "scope": "县域搜救专项",
                "demand_source": "演练指挥协调组任务单QL-RW-0914-01", "demand_total": 500,
                "demand_components": {"县救援力量": 240, "专业力量": 160, "机动力量": 100},
                "available_total": 320, "available_components": {"县救援力量": 160, "专业力量": 100, "机动力量": 60},
                "not_arrived_reserve": 60, "reserve_included_in_available": False,
                "gap_precomputed_in_source": False,
            },
            "rescue_1100": {
                "source": "02-变更参考/05-人员到位更新.md", "as_of": "2026-09-14T11:00:00+08:00",
                "new_arrivals": 80, "new_arrival_components": {"原后备到位": 60, "新增增援": 20},
                "available_total": 400, "demand_total": 500, "original_320_continue_available": True,
                "updated_category_breakdown": "unknown", "changes_only_rescue_arrival_and_available_count": True,
            },
        },
        "expected_inferences_not_in_initial_sources": [
            {"id": "water_dependency_impact", "required_sources": [1, 2], "expected": "云桥体育馆安置点因依赖柳溪供水站而面临供水保障风险；不能据现有材料断言安置点已经耗尽存水或完全无水"},
            {"id": "responsibility_routing", "required_sources": [1, 2, 3], "expected": "由县生活保障组对云桥体育馆安置点核实保障条件并按指引处置，由供水巡查组核实和续报柳溪供水站状态"},
            {"id": "negative_scope", "required_sources": [1, 2], "expected": "城北中学由城北备用水井供水，现有材料未发现其依赖柳溪；不得扩大为无风险结论"},
        ],
        "expected_calculations_not_for_upload": {
            "demand_total": "240 + 160 + 100 = 500",
            "available_0900": "160 + 100 + 60 = 320",
            "gap_0900": "500 - 320 = 180",
            "new_1100": "60 + 20 = 80",
            "available_1100": "320 + 80 = 400",
            "gap_1100": "500 - 400 = 100",
            "gap_change": "100 - 180 = -80",
        },
        "forbidden_assertions": ["把60名未到场后备计入09:00可用人数", "按安置人数推算搜救需求", "11:00仍使用180人缺口", "自行分配新增80人的力量类别", "城北中学没有供水风险", "水质合格或污染已确认", "供水已恢复", "送水车辆已经出发或到达", "将演示材料描述为真实灾情或调度命令"],
        "character_counts_chinese": {
            "01": sum("\u4e00" <= c <= "\u9fff" for c in content_text(FACILITY)),
            "02": sum("\u4e00" <= c <= "\u9fff" for c in content_text(INCIDENT)),
            "03": sum("\u4e00" <= c <= "\u9fff" for c in content_text(GUIDE)),
            "04": sum("\u4e00" <= c <= "\u9fff" for c in RESCUE_MD),
            "05": sum("\u4e00" <= c <= "\u9fff" for c in UPDATE_MD),
        },
    }
    (BASE / "source_truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for name, count in truth["character_counts_chinese"].items():
        assert 500 <= count <= 950, (name, count)
    assert "180" not in RESCUE_MD
    assert "云桥" not in content_text(INCIDENT) and "城北" not in content_text(INCIDENT)
    assert "11:00" not in RESCUE_MD
    assert len(list(INITIAL.iterdir())) == 4, "首轮目录必须只有4份上传材料"
    print(json.dumps({"base": str(BASE), "chinese_characters": truth["character_counts_chinese"], "initial_files": sorted(p.name for p in INITIAL.iterdir())}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
