from __future__ import annotations

import json
import html
import re
import shutil
import subprocess
from functools import lru_cache
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill


CONTENT_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "evidence_docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "json": "application/json",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "geojson": "application/geo+json",
}


@lru_cache(maxsize=1)
def _cjk_font_name() -> str:
    """Select an installed CJK font so server-side PDF rendering keeps Chinese text."""
    executable = shutil.which("fc-list")
    if executable:
        completed = subprocess.run(
            [executable, ":lang=zh", "family"],
            capture_output=True,
            text=True,
            timeout=10,
            check=False,
        )
        available = completed.stdout.casefold()
        for candidate in (
            "Noto Sans CJK SC",
            "WenQuanYi Zen Hei",
            "Arial Unicode MS",
            "PingFang SC",
            "Hiragino Sans GB",
            "SimSun",
        ):
            if candidate.casefold() in available:
                return candidate
    return "Arial Unicode MS"


def _node_text(node: dict[str, Any]) -> str:
    if "text" in node:
        return str(node.get("text") or "")
    return "".join(_node_text(item) for item in (node.get("children") or []) if isinstance(item, dict))


def _clean_citation_text(value: str, max_length: int = 420) -> str:
    """Render retrieved Markdown as a compact, safe source excerpt."""
    cleaned = html.unescape(value)
    cleaned = re.sub(r"(^|[\s：])#{1,6}\s+", r"\1", cleaned, flags=re.MULTILINE)
    cleaned = re.sub(r"^\s{0,3}>\s?", "", cleaned, flags=re.MULTILINE)
    cleaned = re.sub(r"(?:\*\*|__|`)(.*?)(?:\*\*|__|`)", r"\1", cleaned)
    cleaned = re.sub(r"\[(.*?)\]\([^)]*\)", r"\1", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if len(cleaned) > max_length:
        return cleaned[: max_length - 1].rstrip() + "…"
    return cleaned


def _unique_citation_bindings(bindings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Return one evidence row per visible citation number.

    A citation node can appear multiple times in the formal report.  The
    evidence appendix is an index, not an occurrence log, so repeating the
    same source for every inline occurrence makes it noisy and misleading.
    """
    unique: dict[int, dict[str, Any]] = {}
    for item in bindings:
        if item.get("block_type") != "knowledge_citation":
            continue
        metadata = item.get("metadata_json") or {}
        try:
            number = int(metadata.get("citation_number") or 0)
        except (TypeError, ValueError):
            continue
        if number > 0:
            unique.setdefault(number, item)
    return [unique[number] for number in sorted(unique)]


def bindings_for_content(content: list[dict[str, Any]], bindings: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Historical generations can reuse citation numbers, never block IDs.
    Export only bindings present in the chosen immutable document version.
    """
    from .writing import walk_plate_nodes
    visible = {node.get("id") for node in walk_plate_nodes(content) if node.get("id")}
    return {key: binding for key, binding in bindings.items() if key in visible}


def _set_run_font(run, size: float = 11, bold: bool = False) -> None:
    font_name = _cjk_font_name()
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _remove_paragraph_borders(paragraph_or_style) -> None:
    paragraph_properties = paragraph_or_style._element.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def _display_title(title: str) -> str:
    """Remove controlled workspace qualifiers from the formal document title."""
    value = re.sub(r"\s*（知识版本\s*\d+）\s*$", "", title).strip()
    value = re.sub(r"\s*[（(](?:盲测|测试|演示|生成于|\d{4}[-年]\d{1,2})[^）)]*[）)]\s*$", "", value).strip()
    return value


def _display_source_locator(page: Any, structural_path: Any) -> str:
    """Translate parser-internal paths into reviewer-facing source locations."""
    if page not in (None, ""):
        return f"第 {page} 页"
    value = str(structural_path or "").strip()
    paragraph = re.fullmatch(r"paragraphs/(\d+)", value)
    if paragraph:
        return f"正文第 {int(paragraph.group(1)) + 1} 段"
    sheet = re.fullmatch(r"sheets/(.+)", value)
    if sheet:
        return f"工作表“{sheet.group(1)}”"
    if value in {"", "document"}:
        return "文档正文"
    return value


def _source_type_label(value: Any) -> str:
    return {
        "policy_document": "文档来源",
        "official_brief": "官方简报",
        "computation": "确定性测算",
        "semantica_inference": "规则推演",
        "structured_query": "实时数据",
        "manual": "人工确认",
    }.get(str(value or ""), "其他来源")


def _formula_label(value: Any) -> str:
    return {
        "resource_gap": "缺口 = max(需求量 − 可用量, 0)",
    }.get(str(value or ""), str(value or "已登记公式"))


def _configure_docx(document: Document, title: str) -> None:
    section = document.sections[0]
    # Formal Chinese documents use A4. Template-specific overrides can still
    # be applied by the export-template layer before publication.
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.78)
    styles = document.styles
    font_name = _cjk_font_name()
    for style_name, size, bold in (("Normal", 11, False), ("Title", 22, True), ("Heading 1", 16, True), ("Heading 2", 14, True), ("Heading 3", 12, True)):
        style = styles[style_name]
        style.font.name = font_name
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        if style_name == "Title":
            _remove_paragraph_borders(style)
    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _remove_paragraph_borders(title_paragraph)
    _set_run_font(title_paragraph.add_run(title), 22, True)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(header.add_run(title), 9, False)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    _set_run_font(run, 9, False)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
    tail = footer.add_run(" 页")
    _set_run_font(tail, 9, False)


def _add_table(document: Document, node: dict[str, Any]) -> None:
    rows = [item for item in (node.get("children") or []) if isinstance(item, dict)]
    matrix = []
    for row in rows:
        cells = [item for item in (row.get("children") or []) if isinstance(item, dict)]
        matrix.append([_node_text(cell).strip() for cell in cells])
    width = max((len(row) for row in matrix), default=0)
    if not width:
        return
    table = document.add_table(rows=len(matrix), cols=width)
    table.style = "Table Grid"
    for row_index, values in enumerate(matrix):
        for column_index in range(width):
            cell = table.cell(row_index, column_index)
            cell.text = values[column_index] if column_index < len(values) else ""
            cell.vertical_alignment = 1
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    _set_run_font(run, 10, row_index == 0)


def _append_plate_children(paragraph, children: list[dict[str, Any]]) -> None:
    """Render inline Plate nodes, including visible evidence references."""
    for child in children:
        if not isinstance(child, dict):
            continue
        child_type = str(child.get("type") or "")
        if child_type == "knowledge_citation":
            label = str(child.get("citation_label") or "[依据]")
            run = paragraph.add_run(label)
            _set_run_font(run, 9, True)
            run.font.superscript = True
            continue
        if "text" in child:
            run = paragraph.add_run(str(child.get("text") or ""))
            _set_run_font(run, 11, bool(child.get("bold")))
            run.italic = bool(child.get("italic"))
            run.underline = bool(child.get("underline"))
            continue
        _append_plate_children(paragraph, [item for item in child.get("children") or [] if isinstance(item, dict)])


def build_docx(path: Path, *, title: str, content: list[dict[str, Any]], audit_summary: dict[str, Any]) -> None:
    document = Document()
    formal_title = _display_title(title)
    _configure_docx(document, formal_title)
    for index, node in enumerate(content):
        node_type = str(node.get("type") or "p")
        text = _node_text(node).strip()
        if index == 0 and node_type in {"h1", "heading1"} and text in {title, formal_title}:
            continue
        if node_type == "knowledge_citation":
            text = _clean_citation_text(text)
        if node_type == "table":
            _add_table(document, node)
            continue
        if not text:
            continue
        if node_type in {"h1", "heading1"}:
            paragraph = document.add_paragraph(style="Heading 1")
        elif node_type in {"h2", "heading2"}:
            paragraph = document.add_paragraph(style="Heading 2")
        elif node_type in {"h3", "heading3"}:
            paragraph = document.add_paragraph(style="Heading 3")
        elif node_type in {"ul", "bulleted-list"}:
            paragraph = document.add_paragraph(style="List Bullet")
        elif node_type in {"ol", "numbered-list"}:
            paragraph = document.add_paragraph(style="List Number")
        else:
            paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.4
        prefix = ""
        if node_type == "knowledge_citation":
            prefix = "来源依据  "
        elif node_type == "computed_metric":
            prefix = "计算结果  "
        elif node_type == "inference_conclusion":
            prefix = "推演结论  "
        elif node_type == "manual_assumption":
            prefix = "人工假设  "
        if prefix:
            _set_run_font(paragraph.add_run(prefix), 10, True)
        children = [item for item in node.get("children") or [] if isinstance(item, dict)]
        if children and node_type not in {"knowledge_citation", "computed_metric", "inference_conclusion", "manual_assumption"}:
            _append_plate_children(paragraph, children)
        else:
            _set_run_font(paragraph.add_run(text), 11, False)
    document.save(path)


def build_evidence_docx(
    path: Path,
    *,
    title: str,
    facts: list[dict[str, Any]],
    computations: list[dict[str, Any]],
    plans: list[dict[str, Any]],
    audit_summary: dict[str, Any],
    bindings: list[dict[str, Any]],
) -> None:
    """Build a separate, auditable evidence report without polluting the formal report."""
    document = Document()
    _configure_docx(document, f"{_display_title(title)}—生成依据")
    intro = document.add_paragraph()
    _set_run_font(intro.add_run("本文件记录报告使用的知识版本、已核验输入、确定性测算、规则推演和正文来源。它用于核验，不属于正式报告正文。"), 11)

    bound_conclusions = {
        str(ref.get("ref"))
        for binding in bindings
        for ref in (binding.get("metadata_json") or {}).get("knowledge_evidence") or []
        if ref.get("kind") == "inference" and ref.get("ref")
    } | {
        str(binding.get("source_id") or binding.get("block_id"))
        for binding in bindings if binding.get("block_type") == "inference_conclusion"
    }
    summary_rows = [
        ("知识产品版本", str(audit_summary.get("knowledge_product_release") or "-")),
        ("场景包版本", str(audit_summary.get("scenario_package_version") or "-")),
        ("已核验事实", str(audit_summary.get("verified_fact_count") or 0)),
        ("确定性计算", str(audit_summary.get("computation_count") or 0)),
        ("正文规则结论", str(len(bound_conclusions))),
        ("采用方案", str(audit_summary.get("selected_plan") or "未选择")),
    ]
    heading = document.add_paragraph(style="Heading 1")
    _set_run_font(heading.add_run("一、生成基线"), 16, True)
    table = document.add_table(rows=len(summary_rows) + 1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "当前值"
    for index, (label, value) in enumerate(summary_rows, 1):
        table.rows[index].cells[0].text = label
        table.rows[index].cells[1].text = value

    heading = document.add_paragraph(style="Heading 1")
    _set_run_font(heading.add_run("二、已核验输入"), 16, True)
    fact_table = document.add_table(rows=max(1, len(facts)) + 1, cols=5)
    fact_table.style = "Table Grid"
    for index, label in enumerate(("事实", "当前值", "单位", "来源", "版本")):
        fact_table.rows[0].cells[index].text = label
    for row_index, item in enumerate(facts, 1):
        value = item.get("value") or {}
        values = (
            item.get("label"),
            value.get("number", value.get("text", value.get("value"))),
            item.get("unit") or "",
            _source_type_label(item.get("source_type")),
            item.get("version") or "",
        )
        for column_index, value in enumerate(values):
            fact_table.rows[row_index].cells[column_index].text = str(value if value is not None else "")

    heading = document.add_paragraph(style="Heading 1")
    _set_run_font(heading.add_run("三、计算与推演"), 16, True)
    for item in computations:
        result = item.get("result") or {}
        output = result.get("output_fact") or {}
        paragraph = document.add_paragraph(style="List Bullet")
        _set_run_font(paragraph.add_run(f"{output.get('label') or result.get('operation')}：{result.get('value')} {output.get('unit') or ''}；{_formula_label(result.get('operation'))}。"), 11)
    for item in plans:
        if item.get("status") != "selected":
            continue
        route = (item.get("result") or {}).get("route") or {}
        paragraph = document.add_paragraph(style="List Bullet")
        _set_run_font(paragraph.add_run(f"采用方案：{item.get('name')}；路线：{' → '.join(route.get('path') or [])}；预计 {route.get('minutes', '—')} 分钟。"), 11)

    heading = document.add_paragraph(style="Heading 1")
    _set_run_font(heading.add_run("四、正文来源"), 16, True)
    for item in _unique_citation_bindings(bindings):
        metadata = item.get("metadata_json") or {}
        number = metadata.get("citation_number") or "-"
        semantic = metadata.get("knowledge_evidence") or []
        if semantic:
            for reference in semantic:
                kind = "规则结论" if reference.get("kind") == "inference" else "已有关系"
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.keep_with_next = True
                _set_run_font(paragraph.add_run(f"[{number}] {kind}：{reference.get('text') or '请在平台核验'}"), 10, True)
                for index, premise in enumerate(reference.get("premises") or [], 1):
                    source = premise.get("source") or {}
                    location = _display_source_locator(source.get("page_number"), source.get("structural_path"))
                    version_label = f"，版本 {source['version_number']}" if source.get("version_number") else ""
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.left_indent = Pt(10)
                    paragraph.paragraph_format.space_after = Pt(3)
                    _set_run_font(paragraph.add_run(f"前提{index}：{premise.get('text') or '来源关系'}。来源：{source.get('title') or '知识材料'}{version_label}，{location}。"), 10)
            continue
        source = metadata.get("source_title") or "知识材料"
        page = metadata.get("page_number")
        path_value = metadata.get("structural_path")
        location = _display_source_locator(page, path_value)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1.5)
        paragraph.paragraph_format.line_spacing = 1.0
        _set_run_font(paragraph.add_run(f"[{number}] {source}，{location}。"), 10)

    for table_value in document.tables:
        for row_index, row in enumerate(table_value.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_run_font(run, 10, row_index == 0)
    document.save(path)


def build_xlsx(path: Path, *, facts: list[dict[str, Any]], computations: list[dict[str, Any]], plans: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)

    fact_labels = {
        str(item.get("fact_key") or ""): str(item.get("label") or item.get("fact_key") or "")
        for item in facts
    }

    def fact_value(item: dict[str, Any]) -> Any:
        value = item.get("value") or {}
        if "number" in value:
            return value.get("number")
        if "text" in value:
            return value.get("text")
        if "boolean" in value:
            return "是" if value.get("boolean") else "否"
        return value.get("value")

    facts_sheet = workbook.active
    facts_sheet.title = "已核验事实"
    facts_sheet.append(["事实", "值", "单位", "来源", "版本", "状态"])
    for item in facts:
        facts_sheet.append(
            [
                item.get("label"),
                fact_value(item),
                item.get("unit"),
                _source_type_label(str(item.get("source_type") or "")),
                item.get("version"),
                {"verified": "已核验", "unverified": "待核验"}.get(
                    str(item.get("verification_status") or ""),
                    item.get("verification_status"),
                ),
            ]
        )

    computation_sheet = workbook.create_sheet("确定性计算")
    computation_sheet.append(["计算项", "结果", "单位", "公式", "输入事实"])
    for item in computations:
        result = item.get("result") or {}
        output = result.get("output_fact") or {}
        dependency_labels = [
            fact_labels.get(str(key), str(key))
            for key in (result.get("dependencies") or {}).values()
        ]
        computation_sheet.append(
            [
                output.get("label") or result.get("operation"),
                result.get("value"),
                output.get("unit"),
                _formula_label(str(result.get("operation") or "")),
                "、".join(dependency_labels),
            ]
        )

    plan_sheet = workbook.create_sheet("备选方案")
    plan_sheet.append(["方案", "优化目标", "路线", "预计分钟", "路线风险", "状态"])
    for item in plans:
        route = (item.get("result") or {}).get("route") or {}
        plan_sheet.append(
            [
                item.get("name"),
                {"balanced": "综合平衡", "safety": "安全优先", "speed": "速度优先"}.get(
                    str(item.get("objective") or ""), item.get("objective")
                ),
                " → ".join(route.get("path") or []),
                route.get("minutes"),
                route.get("risk"),
                {"selected": "已采用", "candidate": "备选"}.get(
                    str(item.get("status") or ""), item.get("status")
                ),
            ]
        )

    widths = {
        "已核验事实": [24, 30, 12, 16, 10, 12],
        "确定性计算": [24, 16, 10, 34, 40],
        "备选方案": [18, 16, 50, 14, 14, 12],
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = sheet.dimensions
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.45
        sheet.page_margins.bottom = 0.45
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, width in enumerate(widths[sheet.title], 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(path)


def build_geojson(path: Path, *, selected_plan: dict[str, Any], coordinates: dict[str, list[float]]) -> None:
    route = (selected_plan.get("result") or {}).get("route") or {}
    nodes = route.get("path") or []
    missing = [node for node in nodes if node not in coordinates]
    if not nodes or missing:
        raise ValueError("采用方案缺少可核验的路线坐标")
    payload = {
        "type": "FeatureCollection",
        "features": [
            {
                "type": "Feature",
                "properties": {
                    "name": selected_plan.get("name"),
                    "objective": selected_plan.get("objective"),
                    "minutes": route.get("minutes"),
                    "risk": route.get("risk"),
                },
                "geometry": {"type": "LineString", "coordinates": [coordinates[node] for node in nodes]},
            }
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def build_export_artifact(
    path: Path,
    *,
    output_format: str,
    title: str,
    content: list[dict[str, Any]],
    facts: list[dict[str, Any]],
    computations: list[dict[str, Any]],
    plans: list[dict[str, Any]],
    audit_summary: dict[str, Any],
    bindings: list[dict[str, Any]] | None = None,
    coordinates: dict[str, list[float]] | None = None,
) -> Path:
    if output_format == "docx":
        build_docx(path, title=title, content=content, audit_summary=audit_summary)
    elif output_format == "evidence_docx":
        build_evidence_docx(
            path,
            title=title,
            facts=facts,
            computations=computations,
            plans=plans,
            audit_summary=audit_summary,
            bindings=bindings or [],
        )
    elif output_format == "pdf":
        source = path.with_suffix(".docx")
        build_docx(source, title=title, content=content, audit_summary=audit_summary)
        executable = shutil.which("libreoffice") or shutil.which("soffice")
        if executable is None:
            raise RuntimeError("服务器未安装 LibreOffice，无法生成 PDF")
        completed = subprocess.run(
            [executable, "--headless", "--convert-to", "pdf", "--outdir", str(path.parent), str(source)],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        source.unlink(missing_ok=True)
        if completed.returncode != 0 or not path.exists():
            raise RuntimeError("LibreOffice 未能生成 PDF")
    elif output_format == "json":
        payload = {"title": title, "content": content, "facts": facts, "computations": computations, "plans": plans, "audit_summary": audit_summary}
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    elif output_format == "xlsx":
        build_xlsx(path, facts=facts, computations=computations, plans=plans)
    elif output_format == "geojson":
        selected = next((item for item in plans if item.get("status") == "selected"), None)
        if selected is None:
            raise ValueError("请先选择一套备选方案")
        build_geojson(path, selected_plan=selected, coordinates=coordinates or {})
    else:
        raise ValueError("不支持的导出格式")
    if not path.exists() or path.stat().st_size == 0:
        raise RuntimeError("导出文件为空")
    return path
